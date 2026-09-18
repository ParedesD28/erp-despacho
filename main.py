"""Sistema de Gestión Judicial y ERP Inmobiliario.

Módulo principal de producción que consolida todas las rutas operativas,
liquidador de expensas PH, supervisión de agente y administración de expedientes.
"""
from __future__ import annotations

import sms_router
import io
import json
import os
import re
from datetime import date, datetime, timedelta
from typing import Any, List
from urllib.parse import urlencode

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
from psycopg2.extras import RealDictCursor
from fastapi import (
    FastAPI,
    File,
    Form,
    HTTPException,
    Request,
    UploadFile,
    status,
)
from fastapi.encoders import jsonable_encoder
from fastapi.responses import (
    FileResponse,
    JSONResponse,
    RedirectResponse,
    StreamingResponse,
)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

import db
import security
from security import (
    SESSION_COOKIE,
    clear_session_cookie,
    set_session_cookie,
    verify_password,
    verify_session,
)
import observability
from observability import _json_log
import tasas
import liquidador
import exportaciones
import expedientes_service
import catalogos_service
import bot_api
import agent_supervision

app = FastAPI(title="Gestión Judicial ERP", version="2.0.0")
templates = Jinja2Templates(directory="templates")

import api_recaudos

app.include_router(api_recaudos.router)

def render_template(name: str, context: dict, status_code: int = 200):
    """Renderiza plantillas Jinja2 siendo compatible con cualquier versión de Starlette/FastAPI."""
    try:
        return templates.TemplateResponse(name, context, status_code=status_code)
    except TypeError:
        return templates.TemplateResponse(request=context.get("request"), name=name, context=context, status_code=status_code)


# Montaje de archivos estáticos
os.makedirs("static/pdfs", exist_ok=True)
app.mount("/static", StaticFiles(directory="static"), name="static")

# Inclusión de routers modulares M2M y supervisión
app.include_router(bot_api.router)
app.include_router(agent_supervision.router)
app.include_router(sms_router.router)

# Configuración global de observabilidad y captura de excepciones
observability.install_exception_handling(app)


# ==============================================================================
# MIDDLEWARE DE SEGURIDAD PRODUCTIVA
# ==============================================================================
def _is_public_path(path: str) -> bool:
    if path in ("/login", "/health"):
        return True
    if path.startswith("/static/") or path.startswith("/api/bot/"):
        return True
    if path.startswith("/sms/api/"):
        return True
    return False


@app.middleware("http")
async def production_security_middleware(request: Request, call_next):
    path = request.url.path

    # Autenticación segura en /login
    if path == "/login" and request.method == "POST":
        try:
            form = await request.form()
            email = str(form.get("email") or "").strip().lower()
            password = str(form.get("password") or "")
            if not email or not password:
                return render_template("login.html", {"request": request, "error": "Credenciales incorrectas."}, status_code=401)

            conn = db.get_connection()
            try:
                with conn.cursor(cursor_factory=RealDictCursor) as cur:
                    cur.execute(
                        "SELECT id, email, password FROM abogados WHERE LOWER(email)=LOWER(%s) LIMIT 1",
                        (email,),
                    )
                    usuario = cur.fetchone()
            finally:
                conn.release()

            if not usuario or not verify_password(password, usuario.get("password")):
                _json_log("INFO", "login_failed")
                return render_template("login.html", {"request": request, "error": "Credenciales incorrectas."}, status_code=401)

            response = RedirectResponse(url="/dashboard", status_code=303)
            set_session_cookie(response, str(usuario["id"]))
            _json_log("INFO", "login_success")
            return response
        except Exception:
            _json_log("ERROR", "login_error")
            return render_template("login.html", {"request": request, "error": "No fue posible iniciar sesión."}, status_code=500)

    if path == "/logout":
        response = RedirectResponse(url="/login", status_code=303)
        clear_session_cookie(response)
        return response

    if not _is_public_path(path):
        user_id = verify_session(request.cookies.get(SESSION_COOKIE))
        if not user_id:
            response = RedirectResponse(url="/login", status_code=303)
            clear_session_cookie(response)
            return response
        request.state.user_id = user_id

    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    if request.url.scheme == "https":
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response


# ==============================================================================
# HELPERS COMUNES
# ==============================================================================
def _redirect(path: str, **params) -> RedirectResponse:
    if not params:
        return RedirectResponse(url=path, status_code=303)
    sep = "&" if "?" in path else "?"
    query = urlencode(params)
    return RedirectResponse(url=f"{path}{sep}{query}", status_code=303)


def _generar_radicado_interno(cur) -> str:
    """Genera el ID interno secuencial EXP-xxxx, independiente del radicado Rama."""
    cur.execute("SELECT pg_advisory_xact_lock(%s)", (71302541,))
    cur.execute(
        """
        SELECT COALESCE(
            MAX(CAST(SUBSTRING(radicado_interno FROM 5) AS BIGINT)),
            0
        ) AS max
        FROM procesos
        WHERE radicado_interno ~ '^EXP-[0-9]+$'
        """
    )
    row = cur.fetchone()
    # crear_expediente_completo usa RealDictCursor, por lo que la fila
    # debe leerse por nombre de columna, no por índice numérico.
    ultimo = int(row["max"] or 0) if row else 0
    return f"EXP-{ultimo + 1:04d}"

def cargar_inmuebles_ph(conn=None) -> list[dict]:
    external_conn = conn is not None
    if not external_conn:
        conn = db.get_connection()
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT i.id, i.conjunto_residencial, i.torre_apto, c.identificacion, c.nombre,
                   COALESCE(p.radicado_interno, 'SIN EXPEDIENTE') AS expediente
            FROM inmuebles_ph i
            JOIN contactos c ON i.contacto_id = c.id
            LEFT JOIN procesos p ON p.inmueble_id = i.id
            UNION
            SELECT i.id, i.conjunto_residencial, i.torre_apto, c.identificacion, c.nombre,
                   COALESCE(p.radicado_interno, 'SIN EXPEDIENTE') AS expediente
            FROM inmuebles_ph i
            JOIN procesos p ON p.inmueble_id = i.id
            JOIN procesos_litisconsorcio pl ON pl.radicado_interno = p.radicado_interno
            JOIN contactos c ON c.identificacion = pl.identificacion_demandado
            ORDER BY expediente DESC, nombre ASC
        """)
        lista = [
            {"id": r[0], "conjunto_residencial": r[1], "torre_apto": r[2], "apto": r[2], "cedula": r[3], "nombre": r[4], "expediente": r[5]}
            for r in cur.fetchall()
        ]
        cur.close()
        return lista
    except Exception as e:
        print(f"[INMUEBLES] Error cargando inmuebles: {e}", flush=True)
        return []
    finally:
        if not external_conn and conn:
            conn.release()


def sumar_dias_habiles(fecha_inicial: date, dias: int) -> date:
    fecha_actual = fecha_inicial
    dias_agregados = 0
    while dias_agregados < dias:
        fecha_actual += timedelta(days=1)
        if fecha_actual.weekday() < 5:
            dias_agregados += 1
    return fecha_actual


# ==============================================================================
# RUTAS BÁSICAS / VISTAS DE NAVEGACIÓN
# ==============================================================================
@app.get("/", include_in_schema=False)
def root_compat():
    return RedirectResponse(url="/dashboard", status_code=303)


@app.get("/health")
def health():
    # Respuesta inmediata para el health check de Render (evita timeouts de 5s)
    return {"status": "ok", "service": "erp-despacho"}


@app.head("/login", include_in_schema=False)
@app.get("/login")
def vista_login(request: Request):
    return render_template("login.html", {"request": request})


@app.get("/dashboard")
def vista_dashboard(request: Request):
    from observability import log_msg
    import time

    log_msg("🔍 [DASHBOARD]", "Petición recibida, solicitando conexión a BD...")
    t0 = time.perf_counter()
    conn = db.get_connection()
    log_msg("🔌 [DASHBOARD]", "Conexión a BD lista", ms=round((time.perf_counter() - t0) * 1000, 1))
    
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            # 1. Total Procesos Activos
            cur.execute("SELECT COUNT(*) AS total FROM procesos WHERE COALESCE(estado, 'Activo') ILIKE 'Activo'")
            row_p = cur.fetchone()
            total_procesos = row_p["total"] if row_p else 0

            # 2. Total Inmuebles
            cur.execute("SELECT COUNT(*) AS total FROM inmuebles_ph")
            row_i = cur.fetchone()
            total_inmuebles = row_i["total"] if row_i else 0

            # 3. Acuerdos de Pago
            hoy = date.today()
            acuerdos_hoy = []
            acuerdos_vencidos = []
            acuerdos_proximos = []
            monto_acuerdos_vigentes = 0.0

            if expedientes_service._table_exists(cur, "acuerdos_pago"):
                cur.execute("""
                    SELECT a.*, i.conjunto_residencial, i.torre_apto
                    FROM acuerdos_pago a
                    LEFT JOIN inmuebles_ph i ON a.inmueble_id = i.id
                    WHERE a.fecha_compromiso = %s AND a.estado = 'PENDIENTE'
                    ORDER BY a.valor_acordado DESC
                """, (hoy,))
                acuerdos_hoy = [dict(r) for r in cur.fetchall()]

                cur.execute("""
                    SELECT a.*, i.conjunto_residencial, i.torre_apto
                    FROM acuerdos_pago a
                    LEFT JOIN inmuebles_ph i ON a.inmueble_id = i.id
                    WHERE a.fecha_compromiso < %s AND a.estado = 'PENDIENTE'
                    ORDER BY a.fecha_compromiso DESC LIMIT 20
                """, (hoy,))
                acuerdos_vencidos = [dict(r) for r in cur.fetchall()]

                cur.execute("""
                    SELECT a.*, i.conjunto_residencial, i.torre_apto
                    FROM acuerdos_pago a
                    LEFT JOIN inmuebles_ph i ON a.inmueble_id = i.id
                    WHERE a.fecha_compromiso > %s AND a.fecha_compromiso <= %s + INTERVAL '7 days' AND a.estado = 'PENDIENTE'
                    ORDER BY a.fecha_compromiso ASC LIMIT 20
                """, (hoy, hoy))
                acuerdos_proximos = [dict(r) for r in cur.fetchall()]

                cur.execute("SELECT COALESCE(SUM(valor_acordado), 0) AS suma FROM acuerdos_pago WHERE estado = 'PENDIENTE'")
                row_s = cur.fetchone()
                monto_acuerdos_vigentes = float(row_s["suma"]) if row_s else 0.0

            # 4. Términos Judiciales Próximos
            terminos_proximos = []
            if expedientes_service._table_exists(cur, "vencimientos"):
                cur.execute("""
                    SELECT * FROM vencimientos
                    WHERE completado = FALSE AND COALESCE(tipo, 'PROCESAL') = 'PROCESAL'
                    ORDER BY fecha_vencimiento ASC LIMIT 10
                """)
                terminos_proximos = [dict(r) for r in cur.fetchall()]

            log_msg("📊 [DASHBOARD BD]", f"Procesos={total_procesos} | Inmuebles={total_inmuebles} | Acuerdos Hoy={len(acuerdos_hoy)}")

            t_render = time.perf_counter()
            resp = render_template(
                "dashboard.html",
                {
                    "request": request,
                    "total_procesos": total_procesos,
                    "total_inmuebles": total_inmuebles,
                    "acuerdos_hoy": acuerdos_hoy,
                    "acuerdos_vencidos": acuerdos_vencidos,
                    "acuerdos_proximos": acuerdos_proximos,
                    "monto_acuerdos_vigentes": monto_acuerdos_vigentes,
                    "terminos_proximos": terminos_proximos,
                    "hoy": str(hoy),
                },
            )
            log_msg("🎨 [DASHBOARD HTML]", "Plantilla generada exitosamente", ms=round((time.perf_counter() - t_render) * 1000, 1))
            return resp
    finally:
        conn.release()


@app.get("/logout")
def cerrar_sesion():
    resp = RedirectResponse(url="/login", status_code=303)
    clear_session_cookie(resp)
    return resp


# ==============================================================================
# CONTACTOS
# ==============================================================================
@app.get("/contactos")
def contactos(request: Request, buscar_cedula: str | None = None):
    conn = db.get_connection()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cols = expedientes_service._cols(cur, "contactos")
            wanted = ["id", "identificacion", "nombre", "tipo", "telefono", "email", "direccion", "ciudad"]
            select_cols = [c for c in wanted if c in cols]
            if not select_cols:
                select_cols = ["id", "identificacion", "nombre", "tipo", "telefono", "email", "direccion", "ciudad"]
            sql = f"SELECT {', '.join(select_cols)} FROM contactos"
            args = []
            if buscar_cedula:
                term = buscar_cedula.split(" - ", 1)[0].strip()
                sql += " WHERE identificacion ILIKE %s OR nombre ILIKE %s"
                args = [f"%{term}%", f"%{buscar_cedula}%"]
            sql += " ORDER BY nombre ASC LIMIT 500"
            cur.execute(sql, args)
            rows = [dict(r) for r in cur.fetchall()]
            contacto_actual = rows[0] if buscar_cedula and rows else None
            return render_template("contactos.html", {"request": request, "lista_contactos": rows, "contacto_actual": contacto_actual})
    finally:
        conn.release()


@app.post("/contactos/guardar")
def guardar_contacto(
    request: Request,
    identificacion: str = Form(...),
    nombre: str = Form(...),
    tipo: str = Form(...),
    telefono: str = Form(""),
    email: str = Form(""),
    direccion: str = Form(""),
    ciudad: str = Form("PEREIRA"),
    id_contacto: str | None = Form(None),
):
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cols = expedientes_service._cols(cur, "contactos")
                values = {
                    "identificacion": identificacion.strip(),
                    "nombre": nombre.strip(),
                    "tipo": tipo,
                    "telefono": telefono.strip(),
                    "email": email.strip(),
                    "direccion": direccion.strip(),
                    "ciudad": ciudad.strip() or "PEREIRA",
                }
                usable = [c for c in values if c in cols]
                if id_contacto and "id" in cols:
                    sets = ", ".join(f"{c}=%s" for c in usable)
                    cur.execute(f"UPDATE contactos SET {sets} WHERE id=%s", [values[c] for c in usable] + [id_contacto])
                else:
                    names = ", ".join(usable)
                    marks = ", ".join(["%s"] * len(usable))
                    if "identificacion" in cols:
                        cur.execute(
                            f"INSERT INTO contactos ({names}) VALUES ({marks}) "
                            f"ON CONFLICT (identificacion) DO UPDATE SET nombre=EXCLUDED.nombre, tipo=EXCLUDED.tipo",
                            [values[c] for c in usable],
                        )
                    else:
                        cur.execute(f"INSERT INTO contactos ({names}) VALUES ({marks})", [values[c] for c in usable])
        return _redirect("/contactos", mensaje="Contacto+guardado")
    except Exception as exc:
        print(f"[CONTACTOS] Error guardando contacto: {exc}", flush=True)
        return _redirect("/contactos", error="No+fue+posible+guardar+el+contacto")
    finally:
        conn.release()


# ==============================================================================
# RADICACIÓN Y PROCESOS
# ==============================================================================

@app.get("/conjuntos")
def conjuntos(request: Request, id: int | None = None):
    conn = db.get_connection()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            conjuntos = catalogos_service.listar_conjuntos(cur, activos=True)
            conjunto_actual = catalogos_service.obtener_conjunto(cur, id) if id else None
        return render_template(
            "conjuntos.html",
            {"request": request, "conjuntos": conjuntos, "conjunto_actual": conjunto_actual},
        )
    finally:
        conn.release()


@app.post("/conjuntos/guardar")
def guardar_conjunto(
    request: Request,
    nombre: str = Form(...),
    nit: str = Form(...),
    telefono: str = Form(""),
    email: str = Form(""),
    direccion: str = Form(""),
    ciudad: str = Form("PEREIRA"),
    id: str | None = Form(None),
):
    nombre = nombre.strip()
    nit = nit.strip()
    if not nombre or not nit:
        return _redirect("/conjuntos", error="Nombre+y+NIT+son+obligatorios")

    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                if id:
                    cur.execute(
                        "SELECT contacto_id FROM conjuntos_residenciales WHERE id=%s FOR UPDATE",
                        (int(id),),
                    )
                    actual = cur.fetchone()
                    if not actual:
                        raise ValueError("El conjunto no existe")
                    contacto_id = actual[0]
                else:
                    contacto_id = None

                cur.execute(
                    """
                    INSERT INTO contactos
                        (identificacion,nombre,tipo,telefono,email,direccion,ciudad)
                    VALUES (%s,%s,'Cliente',%s,%s,%s,%s)
                    ON CONFLICT (identificacion)
                    DO UPDATE SET
                        nombre=EXCLUDED.nombre,
                        tipo='Cliente',
                        telefono=EXCLUDED.telefono,
                        email=EXCLUDED.email,
                        direccion=EXCLUDED.direccion,
                        ciudad=EXCLUDED.ciudad
                    RETURNING id
                    """,
                    (
                        nit,
                        nombre,
                        telefono.strip(),
                        email.strip(),
                        direccion.strip(),
                        ciudad.strip().upper() or "PEREIRA",
                    ),
                )
                contacto_id = cur.fetchone()[0]

                if id:
                    cur.execute(
                        """
                        UPDATE conjuntos_residenciales
                        SET nombre=%s, contacto_id=%s, activo=TRUE
                        WHERE id=%s
                        """,
                        (nombre, contacto_id, int(id)),
                    )
                else:
                    cur.execute(
                        """
                        INSERT INTO conjuntos_residenciales (nombre,contacto_id,activo)
                        VALUES (%s,%s,TRUE)
                        ON CONFLICT (nombre)
                        DO UPDATE SET contacto_id=EXCLUDED.contacto_id,activo=TRUE
                        """,
                        (nombre, contacto_id),
                    )

        return _redirect("/conjuntos", mensaje="Conjunto+guardado")
    except Exception as exc:
        print(f"[CONJUNTOS] Error guardando conjunto: {exc!r}", flush=True)
        return _redirect("/conjuntos", error="No+fue+posible+guardar+el+conjunto")
    finally:
        conn.release()


@app.get("/procesos")
def procesos(request: Request):
    conn = db.get_connection()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT identificacion, nombre FROM contactos WHERE tipo='Cliente' ORDER BY nombre LIMIT 500")
            clientes = [dict(r) for r in cur.fetchall()]
            cur.execute("SELECT identificacion, nombre FROM contactos WHERE tipo='Contraparte' ORDER BY nombre LIMIT 500")
            contrapartes = [dict(r) for r in cur.fetchall()]
            cur.execute("SELECT id, nombre FROM abogados ORDER BY nombre")
            abogados = [dict(r) for r in cur.fetchall()]
            tipos_proceso = catalogos_service.listar_tipos_proceso(cur, activos=True)
            conjuntos = catalogos_service.listar_conjuntos(cur, activos=True)
        return render_template("procesos.html", {
            "request": request,
            "contactos_clientes": clientes,
            "contactos_contrapartes": contrapartes,
            "abogados": abogados,
            "tipos_proceso": tipos_proceso,
            "conjuntos": conjuntos,
        })
    finally:
        conn.release()


@app.post("/crear_expediente_completo")
async def crear_expediente_completo(request: Request):
    form = await request.form()
    tipo_proceso_codigo = str(form.get("tipo_proceso") or "").strip().upper()
    conjunto_id_raw = str(form.get("conjunto_id") or "").strip()
    conjunto_nombre = str(form.get("conjunto_residencial") or "").strip()
    tipo_cartera_form = str(form.get("tipo_cartera") or "").strip().upper()
    radicado_rama = str(form.get("radicado_rama") or "").strip()
    naturaleza = str(form.get("naturaleza", "")).strip()
    juzgado = f"{form.get('juzgado_numero','')} {form.get('juzgado_tipo','')} - {form.get('juzgado_ciudad','')}".strip()
    apto = str(form.get("apto", "")).strip()
    pretensiones = str(form.get("pretensiones", "0")).strip() or "0"
    abogado_id = str(form.get("abogado_id", "")).strip() or None
    medidas = str(form.get("medidas_cautelares", "")).strip()

    def split_values(name, new_id, new_name):
        vals = [x.strip() for x in form.getlist(name) if str(x).strip()]
        ids = [x.strip() for x in form.getlist(new_id) if str(x).strip()]
        names = [x.strip() for x in form.getlist(new_name) if str(x).strip()]
        return vals, list(zip(ids, names))

    demandantes, nuevos_dem = split_values("demandantes_existentes", "nuevo_dem_id", "nuevo_dem_nombre")
    demandados, nuevos_ddo = split_values("demandados_existentes", "nuevo_ddo_id", "nuevo_ddo_nombre")

    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                # Si el usuario no indica tipo, mantenemos compatibilidad con el flujo
                # anterior: un expediente con conjunto es PH; sin conjunto queda como OTRO.
                if not tipo_proceso_codigo:
                    tipo_proceso_codigo = "CUOTAS_ADMINISTRACION" if (conjunto_id_raw or conjunto_nombre) else "OTRO"

                tipo_proceso = catalogos_service.obtener_tipo_proceso(cur, tipo_proceso_codigo)
                if not tipo_proceso:
                    raise ValueError("Tipo de proceso no válido")

                tipo_cartera = tipo_cartera_form or str(tipo_proceso["tipo_cartera_default"]).upper()
                if tipo_cartera not in {"JURIDICO", "PREJURIDICO"}:
                    raise ValueError("Tipo de cartera no válido")

                conjunto = None
                if tipo_proceso["requiere_conjunto"]:
                    if conjunto_id_raw:
                        conjunto = catalogos_service.obtener_conjunto(cur, int(conjunto_id_raw))
                    elif conjunto_nombre:
                        cur.execute(
                            """
                            SELECT id
                            FROM conjuntos_residenciales
                            WHERE activo=TRUE AND UPPER(BTRIM(nombre))=UPPER(BTRIM(%s))
                            LIMIT 1
                            """,
                            (conjunto_nombre,),
                        )
                        row = cur.fetchone()
                        if row:
                            conjunto = catalogos_service.obtener_conjunto(cur, int(row["id"]))
                    if not conjunto:
                        raise ValueError("Debes seleccionar un conjunto residencial válido")
                    if not conjunto.get("contacto_id"):
                        raise ValueError("El conjunto debe tener configurada su persona jurídica antes de radicar cuotas de administración")

                    # Para cuotas de administración el demandante no se digita:
                    # sale automáticamente de la persona jurídica del conjunto.
                    cur.execute(
                        "SELECT identificacion FROM contactos WHERE id=%s",
                        (conjunto["contacto_id"],),
                    )
                    row = cur.fetchone()
                    if not row or not row["identificacion"]:
                        raise ValueError("La persona jurídica del conjunto no tiene identificación configurada")
                    demandantes = [str(row["identificacion"])]

                if tipo_proceso["requiere_inmueble"] and not apto:
                    raise ValueError("Debes indicar la torre/apartamento del inmueble")

                if tipo_proceso["requiere_juzgado"] and tipo_cartera == "JURIDICO":
                    if not juzgado or juzgado.startswith("-"):
                        raise ValueError("El juzgado es obligatorio para este tipo de proceso jurídico")

                if tipo_proceso["requiere_documento"] and not str(form.get("documento_referencia") or "").strip():
                    raise ValueError("Debes indicar el número o referencia del documento título")

                for ident, nombre in nuevos_dem:
                    cur.execute(
                        """
                        INSERT INTO contactos (identificacion,nombre,tipo,ciudad)
                        VALUES (%s,%s,'Cliente','PEREIRA')
                        ON CONFLICT (identificacion) DO UPDATE SET nombre=EXCLUDED.nombre
                        """,
                        (ident, nombre),
                    )
                for ident, nombre in nuevos_ddo:
                    cur.execute(
                        """
                        INSERT INTO contactos (identificacion,nombre,tipo,ciudad)
                        VALUES (%s,%s,'Contraparte','PEREIRA')
                        ON CONFLICT (identificacion) DO UPDATE SET nombre=EXCLUDED.nombre
                        """,
                        (ident, nombre),
                    )

                demandantes += [x[0] for x in nuevos_dem]
                demandados += [x[0] for x in nuevos_ddo]
                if not demandantes or not demandados:
                    raise ValueError("Debe existir al menos un demandante y un demandado")

                if tipo_cartera == "JURIDICO":
                    radicado_rama = radicado_rama or "EN REPARTO"
                elif not radicado_rama or radicado_rama.upper() == "EN REPARTO":
                    radicado_rama = "PREJURIDICO"

                radicado_interno = _generar_radicado_interno(cur)

                if tipo_cartera == "JURIDICO" and radicado_rama.upper() != "EN REPARTO":
                    cur.execute(
                        "SELECT 1 FROM procesos WHERE radicado_rama=%s LIMIT 1",
                        (radicado_rama,),
                    )
                    if cur.fetchone():
                        raise ValueError("El radicado Rama Judicial ya existe")

                all_ids = list(dict.fromkeys(demandantes + demandados))
                placeholders = ",".join(["%s"] * len(all_ids))
                cur.execute(
                    f"SELECT identificacion,nombre FROM contactos WHERE identificacion IN ({placeholders})",
                    all_ids,
                )
                nombres = {str(r["identificacion"]): str(r["nombre"] or "") for r in cur.fetchall()}
                faltantes = [ident for ident in all_ids if ident not in nombres]
                if faltantes:
                    raise RuntimeError("Hay partes seleccionadas que no existen en Contactos")

                inmueble_id = None
                if tipo_proceso["requiere_inmueble"]:
                    conjunto_id = int(conjunto["id"])
                    cur.execute(
                        """
                        SELECT id, contacto_id
                        FROM inmuebles_ph
                        WHERE conjunto_id=%s AND torre_apto=%s
                        LIMIT 1
                        """,
                        (conjunto_id, apto),
                    )
                    inmueble = cur.fetchone()
                    if inmueble:
                        inmueble_id = inmueble["id"]
                    else:
                        cur.execute(
                            """
                            SELECT id FROM contactos
                            WHERE identificacion=%s
                            LIMIT 1
                            """,
                            (demandados[0],),
                        )
                        ddo = cur.fetchone()
                        if not ddo:
                            raise RuntimeError("Demandado no existe")
                        cur.execute(
                            """
                            INSERT INTO inmuebles_ph
                                (contacto_id,conjunto_residencial,conjunto_id,torre_apto)
                            VALUES (%s,%s,%s,%s)
                            RETURNING id
                            """,
                            (ddo["id"], conjunto["nombre"], conjunto_id, apto),
                        )
                        inmueble_id = cur.fetchone()["id"]

                cols = expedientes_service._cols(cur, "procesos")
                data = {
                    "radicado_interno": radicado_interno,
                    "radicado_rama": radicado_rama,
                    "tipo_cartera": tipo_cartera,
                    "tipo_proceso_id": tipo_proceso["id"],
                    "naturaleza": naturaleza or "EJECUTIVO SINGULAR",
                    "juzgado": juzgado if tipo_cartera == "JURIDICO" else "",
                    "estado": "Activo",
                    "id_demandado": " | ".join(demandados),
                    "demandado": " | ".join(nombres.get(x, x) for x in demandados),
                    "inmueble_id": inmueble_id,
                    "id_cliente": " | ".join(demandantes),
                    "demandante": " | ".join(nombres.get(x, x) for x in demandantes),
                    "pretensiones": pretensiones,
                    "medidas_cautelares": medidas,
                    "abogado_id": abogado_id,
                }
                usable = [c for c in data if c in cols and data[c] is not None]
                cur.execute(
                    f"INSERT INTO procesos ({', '.join(usable)}) VALUES ({', '.join(['%s']*len(usable))})",
                    [data[c] for c in usable],
                )

                if tipo_proceso["fuente_saldo"] == "OBLIGACION":
                    documento_referencia = str(form.get("documento_referencia") or "").strip()
                    cur.execute(
                        """
                        INSERT INTO obligaciones
                            (identificacion_deudor,tipo_titulo,numero_documento,capital,fecha_exigibilidad,estado,proceso_id,tipo_proceso_id)
                        VALUES (%s,%s,%s,%s,NULLIF(%s,''),'En Mora',%s,%s)
                        """,
                        (
                            demandados[0],
                            tipo_proceso["codigo"],
                            documento_referencia,
                            pretensiones,
                            str(form.get("fecha_exigibilidad") or "").strip(),
                            radicado_interno,
                            tipo_proceso["id"],
                        ),
                    )

                if expedientes_service._table_exists(cur, "procesos_litisconsorcio"):
                    lit_cols = expedientes_service._cols(cur, "procesos_litisconsorcio")
                    for ident in demandados:
                        payload = {"radicado_interno": radicado_interno, "identificacion_demandado": ident}
                        use = [c for c in payload if c in lit_cols]
                        cur.execute(
                            f"INSERT INTO procesos_litisconsorcio ({', '.join(use)}) VALUES ({', '.join(['%s']*len(use))})",
                            [payload[c] for c in use],
                        )

                if expedientes_service._table_exists(cur, "actuaciones") and tipo_cartera == "JURIDICO":
                    act_cols = expedientes_service._cols(cur, "actuaciones")
                    payload = {
                        "radicado_interno": radicado_interno,
                        "fecha": date.today(),
                        "etapa": "Inicio",
                        "descripcion": "Presentación inicial de la demanda",
                        "usuario": "Sistema",
                        "tipificacion_sugerida": "Radicación",
                    }
                    use = [c for c in payload if c in act_cols]
                    cur.execute(
                        f"INSERT INTO actuaciones ({', '.join(use)}) VALUES ({', '.join(['%s']*len(use))})",
                        [payload[c] for c in use],
                    )

                print(
                    f"[RADICACION] ID interno={radicado_interno} | "
                    f"tipo_proceso={tipo_proceso['codigo']} | "
                    f"radicado_rama={radicado_rama} | tipo_cartera={tipo_cartera}",
                    flush=True,
                )

        return _redirect("/expedientes", mensaje="Proceso+creado+exitosamente")
    except Exception as exc:
        print(f"[PROCESOS] Error creando expediente: {exc!r}", flush=True)
        return _redirect("/procesos", error=str(exc).replace(" ", "+"))
    finally:
        conn.release()


@app.post("/crear_proceso_cascada")
async def crear_proceso_cascada(
    request: Request,
    radicado_interno: str = Form(...),
    radicado_rama: str = Form(...),
    naturaleza: str = Form(...),
    juzgado: str = Form(...),
    demandante_nombre: str = Form(...),
    demandante_cedula: str = Form(...),
    demandado_cedulas: str = Form(...),
    demandado_nombres: str = Form(...),
    conjunto_residencial: str = Form(...),
    nomenclatura_apto: str = Form(...),
):
    radicado_interno = radicado_interno.strip()
    cedulas_list = [c.strip() for c in demandado_cedulas.split("|") if c.strip()]
    nombres_list = [n.strip() for n in demandado_nombres.split("|") if n.strip()]

    if not radicado_interno or not cedulas_list or len(cedulas_list) != len(nombres_list):
        return RedirectResponse(url="/expedientes?error=Datos+de+demandados+invalidos", status_code=303)
    if any(not re.match(r"^\d{6,15}$", c) for c in cedulas_list):
        return RedirectResponse(url="/expedientes?error=Identificacion+invalida", status_code=303)

    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute("SELECT 1 FROM procesos WHERE radicado_interno=%s LIMIT 1", (radicado_interno,))
                if cur.fetchone():
                    return RedirectResponse(url="/expedientes?error=El+radicado+ya+existe", status_code=303)

                deudor_principal_cedula = cedulas_list[0]
                deudor_principal_nombre = nombres_list[0]
                cur.execute("""
                    INSERT INTO contactos (identificacion, nombre, tipo, ciudad)
                    VALUES (%s, %s, 'Cliente', 'PEREIRA')
                    ON CONFLICT (identificacion) DO NOTHING;
                """, (demandante_cedula.strip(), demandante_nombre.strip()))
                cur.execute("""
                    INSERT INTO contactos (identificacion, nombre, tipo, ciudad)
                    VALUES (%s, %s, 'Contraparte', 'PEREIRA')
                    ON CONFLICT (identificacion) DO UPDATE SET nombre = EXCLUDED.nombre
                    RETURNING id;
                """, (deudor_principal_cedula, deudor_principal_nombre))
                res_contacto = cur.fetchone()
                if res_contacto:
                    contacto_id = res_contacto[0]
                else:
                    cur.execute("SELECT id FROM contactos WHERE identificacion = %s", (deudor_principal_cedula,))
                    contacto_id = cur.fetchone()[0]

                cur.execute("""
                    INSERT INTO inmuebles_ph (contacto_id, conjunto_residencial, torre_apto)
                    VALUES (%s, %s, %s)
                    RETURNING id;
                """, (contacto_id, conjunto_residencial.strip(), nomenclatura_apto.strip()))
                inmueble_id = cur.fetchone()[0]

                cur.execute("""
                    INSERT INTO procesos (radicado_interno, radicado_rama, naturaleza, juzgado,
                                          estado, id_demandado, demandado, inmueble_id)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                """, (radicado_interno, radicado_rama.strip(), naturaleza.strip(), juzgado.strip(),
                      'Activo', demandado_cedulas.strip(), demandado_nombres.strip(), inmueble_id))

                if expedientes_service._table_exists(cur, "actuaciones"):
                    cur.execute("""
                        INSERT INTO actuaciones (radicado_interno, fecha, etapa, descripcion, usuario, tipificacion_sugerida)
                        VALUES (%s, CURRENT_DATE, 'Inicio', 'Presentación inicial de la demanda', 'Sistema', 'Radicación')
                    """, (radicado_interno,))
        return RedirectResponse(url="/expedientes?mensaje=Proceso+creado+exitosamente", status_code=303)
    except Exception as e:
        print(f"[CASCADA] Error en la creación: {e}", flush=True)
        return RedirectResponse(url="/expedientes?error=Fallo+la+creacion", status_code=303)
    finally:
        conn.release()


# ==============================================================================
# EXPEDIENTES Y DETALLE
# ==============================================================================
@app.get("/expedientes")
def ver_expedientes(request: Request):
    procesos_lista = expedientes_service.cargar_procesos_general_sin_duplicados()
    return render_template("expedientes.html", {"request": request, "procesos": procesos_lista})


@app.get("/expediente/{radicado}", include_in_schema=False)
def detalle_expediente(request: Request, radicado: str):
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                proceso = expedientes_service._get_process(cur, radicado)
                if not proceso:
                    raise HTTPException(status_code=404, detail="Expediente no encontrado")
                stage = expedientes_service._sync_stage(cur, radicado)
                if stage:
                    proceso["etapa_actual"] = stage
                demandantes = expedientes_service._get_demandantes(cur, proceso)
                demandados = expedientes_service._get_demandados(cur, radicado)
                actuaciones = expedientes_service._get_actuaciones(cur, radicado)
                contactos_opts = expedientes_service._contact_options(cur)
                abogados_opts = expedientes_service._get_abogados(cur)
                ids = [x.get("identificacion") for x in demandantes + demandados]
                acuerdos = expedientes_service._get_crm_agreements(cur, proceso.get("inmueble_id"), ids)
                audit = expedientes_service._audit(cur, radicado)
                return render_template("detalle_expediente_v4.html", {"request": request, "proceso": proceso, "demandantes": demandantes, "demandados": demandados, "contactos": contactos_opts, "abogados": abogados_opts, "actuaciones": actuaciones, "acuerdos_crm": acuerdos, "audit_ediciones": audit, "demandante_ids": {str(x.get("identificacion")) for x in demandantes if x.get("identificacion")}, "demandado_ids": {str(x.get("identificacion")) for x in demandados if x.get("identificacion")}})
    finally:
        conn.release()


@app.post("/expediente/guardar-estructurado", include_in_schema=False)
async def guardar_expediente_estructurado(request: Request):
    form = await request.form()
    radicado = str(form.get("radicado_interno") or "").strip()
    if not radicado:
        return _redirect("/expedientes", error="Expediente+sin+radicado+interno")

    demandante_ids = list(dict.fromkeys(str(x).strip() for x in form.getlist("demandante_id") if str(x).strip()))
    demandado_ids = list(dict.fromkeys(str(x).strip() for x in form.getlist("demandado_id") if str(x).strip()))
    medidas = list(dict.fromkeys(str(x).strip() for x in form.getlist("medida") if str(x).strip()))

    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                proceso = expedientes_service._get_process(cur, radicado)
                if not proceso:
                    raise HTTPException(status_code=404, detail="Expediente no encontrado")
                cols = expedientes_service._cols(cur, "procesos")
                if not expedientes_service._table_exists(cur, "contactos"):
                    raise ValueError("No existe la tabla de contactos")
                if not demandante_ids:
                    raise ValueError("Debe existir al menos un demandante")
                if not demandado_ids:
                    raise ValueError("Debe existir al menos un demandado")

                all_ids = list(dict.fromkeys(demandante_ids + demandado_ids))
                placeholders = ",".join(["%s"] * len(all_ids))
                cur.execute(
                    f"SELECT identificacion, nombre FROM contactos WHERE identificacion IN ({placeholders})",
                    all_ids,
                )
                contacts = {str(r["identificacion"]): r["nombre"] for r in cur.fetchall()}
                missing = [x for x in all_ids if x not in contacts]
                if missing:
                    raise ValueError("Hay una parte seleccionada que no existe en Contactos")

                before = {
                    "proceso": jsonable_encoder(proceso),
                    "demandantes": [dict(x) for x in expedientes_service._get_demandantes(cur, proceso)],
                    "demandados": [dict(x) for x in expedientes_service._get_demandados(cur, radicado)],
                }

                tipo_cartera_form = str(form.get("tipo_cartera") or proceso.get("tipo_cartera") or "JURIDICO").strip().upper()
                rama = str(form.get("radicado_rama") or "").strip()
                if tipo_cartera_form == "JURIDICO" and not rama:
                    rama = "EN REPARTO"
                if tipo_cartera_form != "JURIDICO" and not rama:
                    rama = proceso.get("radicado_rama") or None
                if "radicado_rama" in cols and rama and rama.upper() != "EN REPARTO":
                    cur.execute(
                        "SELECT 1 FROM procesos WHERE radicado_rama=%s AND radicado_interno<>%s LIMIT 1",
                        (rama, radicado),
                    )
                    if cur.fetchone():
                        raise ValueError("El radicado Rama ya pertenece a otro expediente")

                naturaleza = str(form.get("naturaleza") or "").strip()
                juzgado = str(form.get("juzgado") or "").strip()
                if not naturaleza:
                    raise ValueError("La naturaleza es obligatoria")

                if "pretensiones" in form:
                    pretensiones = expedientes_service._parse_money(form.get("pretensiones"))
                else:
                    pretensiones = proceso.get("pretensiones")

                editable = {
                    "radicado_rama": rama,
                    "tipo_cartera": tipo_cartera_form,
                    "naturaleza": naturaleza,
                    "juzgado": juzgado,
                    "pretensiones": pretensiones,
                }
                if "medidas_cautelares" in cols:
                    editable["medidas_cautelares"] = "\n".join(medidas)
                if "abogado_id" in cols:
                    abogado = str(form.get("abogado_id") or "").strip()
                    editable["abogado_id"] = int(abogado) if abogado.isdigit() else None
                if "id_cliente" in cols:
                    editable["id_cliente"] = " | ".join(demandante_ids)
                if "demandante" in cols:
                    editable["demandante"] = " | ".join(contacts[x] for x in demandante_ids)
                if "id_demandado" in cols:
                    editable["id_demandado"] = " | ".join(demandado_ids)
                if "demandado" in cols:
                    editable["demandado"] = " | ".join(contacts[x] for x in demandado_ids)

                usable = [k for k in editable if k in cols and k != "radicado_interno"]
                if usable:
                    cur.execute(
                        f"UPDATE procesos SET {', '.join(f'{k}=%s' for k in usable)} WHERE radicado_interno=%s",
                        [editable[k] for k in usable] + [radicado],
                    )

                if expedientes_service._table_exists(cur, "procesos_litisconsorcio"):
                    lcols = expedientes_service._cols(cur, "procesos_litisconsorcio")
                    if "radicado_interno" in lcols and "identificacion_demandado" in lcols:
                        cur.execute("DELETE FROM procesos_litisconsorcio WHERE radicado_interno=%s", (radicado,))
                        for ident in demandado_ids:
                            cur.execute(
                                "INSERT INTO procesos_litisconsorcio (radicado_interno,identificacion_demandado) VALUES (%s,%s)",
                                (radicado, ident),
                            )

                expedientes_service._ensure_audit_table(cur)
                after_process = expedientes_service._get_process(cur, radicado)
                after = {
                    "proceso": jsonable_encoder(after_process),
                    "demandantes": [dict(x) for x in expedientes_service._get_demandantes(cur, after_process)],
                    "demandados": [dict(x) for x in expedientes_service._get_demandados(cur, radicado)],
                }
                usuario = str(getattr(request.state, "user_id", None) or "ERP")
                cur.execute(
                    "INSERT INTO expediente_ediciones (radicado_interno,usuario,accion,antes,despues) VALUES (%s,%s,%s,%s::jsonb,%s::jsonb)",
                    (radicado, usuario, "EDICION_RADICACION", json.dumps(before, default=str), json.dumps(after, default=str)),
                )
                expedientes_service._sync_stage(cur, radicado)
        return _redirect(f"/expediente/{radicado}", mensaje="Expediente+actualizado")
    except HTTPException:
        raise
    except Exception as exc:
        print(f"[EXPEDIENTE][EDICION] Error {radicado}: {exc!r}", flush=True)
        return _redirect(f"/expediente/{radicado}", error="No+fue+posible+actualizar+el+expediente")
    finally:
        conn.release()


@app.post("/actuacion/nueva")
async def nueva_actuacion(request: Request):
    form = await request.form()
    radicado = str(form.get("radicado_interno") or "").strip()
    if not radicado:
        return _redirect("/expedientes", error="Radicado+interno+requerido")
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                if not expedientes_service._table_exists(cur, "actuaciones"):
                    raise RuntimeError("No existe la tabla actuaciones")
                cols = expedientes_service._cols(cur, "actuaciones")
                payload = {
                    "radicado_interno": radicado,
                    "fecha": form.get("fecha") or date.today(),
                    "etapa": str(form.get("etapa") or "Auto de Trámite / General").strip(),
                    "descripcion": str(form.get("descripcion") or form.get("sub_etapa") or "Actuación registrada").strip(),
                    "usuario": "ERP",
                    "tipificacion_sugerida": str(form.get("sub_etapa") or "Observación").strip(),
                }
                use = [c for c in payload if c in cols]
                cur.execute(
                    f"INSERT INTO actuaciones ({', '.join(use)}) VALUES ({', '.join(['%s']*len(use))}) RETURNING id",
                    [payload[c] for c in use],
                )
                expedientes_service._sync_stage(cur, radicado)
        return _redirect(f"/expediente/{radicado}", mensaje="Actuación+registrada+y+etapa+actualizada")
    except Exception as exc:
        print(f"[ACTUACION] Error registrando {radicado}: {exc!r}", flush=True)
        return _redirect(f"/expediente/{radicado}", error="No+fue+posible+registrar+la+actuación")
    finally:
        conn.release()


@app.post("/actuacion/eliminar")
def eliminar_actuacion(actuacion_id: int, radicado_interno: str):
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute("DELETE FROM actuaciones WHERE id=%s AND radicado_interno=%s", (actuacion_id, radicado_interno))
                if cur.rowcount == 0:
                    raise ValueError("Actuación no encontrada")
                expedientes_service._sync_stage(cur, radicado_interno)
        return _redirect(f"/expediente/{radicado_interno}", mensaje="Actuación+eliminada+y+etapa+recalculada")
    except Exception as exc:
        print(f"[ACTUACION] Error eliminando {actuacion_id}: {exc!r}", flush=True)
        return _redirect(f"/expediente/{radicado_interno}", error="No+fue+posible+eliminar+la+actuación")
    finally:
        conn.release()


# ==============================================================================
# CRM (GESTIÓN EXTRAJUDICIAL Y ANULACIÓN AUDITABLE)
# ==============================================================================
def _ensure_crm_and_vencimientos_schema():
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS gestiones_crm (
                        id BIGSERIAL PRIMARY KEY,
                        inmueble_id INTEGER,
                        identificacion_deudor TEXT,
                        tipo_contacto TEXT,
                        resumen TEXT NOT NULL,
                        promesa_pago_fecha DATE,
                        fecha TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
                        usuario TEXT NOT NULL DEFAULT 'ERP',
                        anulado BOOLEAN NOT NULL DEFAULT FALSE,
                        estado TEXT NOT NULL DEFAULT 'ACTIVO'
                    )
                """)
                cur.execute("ALTER TABLE gestiones_crm ALTER COLUMN inmueble_id DROP NOT NULL")
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS vencimientos (
                        id BIGSERIAL PRIMARY KEY,
                        radicado_interno TEXT NOT NULL,
                        titulo TEXT NOT NULL,
                        fecha_vencimiento DATE NOT NULL,
                        observaciones TEXT,
                        completado BOOLEAN NOT NULL DEFAULT FALSE,
                        created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
                    )
                """)
                cur.execute("ALTER TABLE vencimientos ADD COLUMN IF NOT EXISTS completado BOOLEAN NOT NULL DEFAULT FALSE")
                cur.execute("ALTER TABLE vencimientos ADD COLUMN IF NOT EXISTS created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP")
                cur.execute("ALTER TABLE vencimientos ADD COLUMN IF NOT EXISTS tipo TEXT DEFAULT 'PROCESAL'")
                cur.execute("ALTER TABLE vencimientos ADD COLUMN IF NOT EXISTS valor NUMERIC(14,2) DEFAULT 0")
                cur.execute("ALTER TABLE vencimientos ADD COLUMN IF NOT EXISTS inmueble_id INTEGER")
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS acuerdos_pago (
                        id BIGSERIAL PRIMARY KEY,
                        inmueble_id INTEGER,
                        identificacion_deudor TEXT NOT NULL,
                        nombre_deudor TEXT,
                        telefono TEXT,
                        valor_acordado NUMERIC(14,2) NOT NULL DEFAULT 0,
                        numero_cuotas INTEGER NOT NULL DEFAULT 1,
                        cuota_actual INTEGER NOT NULL DEFAULT 1,
                        fecha_compromiso DATE NOT NULL,
                        estado TEXT NOT NULL DEFAULT 'PENDIENTE',
                        origen TEXT NOT NULL DEFAULT 'ROBOT_IA',
                        observaciones TEXT,
                        recordatorio_previo_enviado BOOLEAN NOT NULL DEFAULT FALSE,
                        recordatorio_dia_enviado BOOLEAN NOT NULL DEFAULT FALSE,
                        recordatorio_mora_enviado BOOLEAN NOT NULL DEFAULT FALSE,
                        created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
                        updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
                    )
                """)
                cur.execute("CREATE INDEX IF NOT EXISTS idx_vencimientos_fecha ON vencimientos (fecha_vencimiento, completado)")
                cur.execute("CREATE INDEX IF NOT EXISTS idx_gestiones_crm_inmueble ON gestiones_crm (inmueble_id, fecha DESC)")
                cur.execute("CREATE INDEX IF NOT EXISTS idx_gestiones_crm_identificacion ON gestiones_crm (identificacion_deudor, fecha DESC)")
                cur.execute("CREATE INDEX IF NOT EXISTS idx_acuerdos_fecha ON acuerdos_pago (fecha_compromiso, estado)")
                cur.execute("CREATE INDEX IF NOT EXISTS idx_acuerdos_deudor ON acuerdos_pago (identificacion_deudor)")
                cur.execute("CREATE INDEX IF NOT EXISTS idx_acuerdos_inmueble ON acuerdos_pago (inmueble_id)")
    except Exception as exc:
        print(f"[SCHEMA] Error asegurando tablas auxiliares: {exc!r}", flush=True)
    finally:
        conn.release()


@app.on_event("startup")
def startup_schema_init():
    """Garantiza la creación y actualización de tablas e índices en Neon automáticamente al iniciar."""
    print("[STARTUP] Verificando y asegurando esquema de base de datos...", flush=True)
    _ensure_crm_and_vencimientos_schema()
    print("[STARTUP] Esquema verificado y asegurado correctamente.", flush=True)


@app.get("/crm")
def crm(
    request: Request,
    buscar_acreedor: str | None = None,
    acreedor_id: str | None = None,
    radicado_interno: str | None = None,
):
    conn = db.get_connection()
    try:
        acreedores = []
        cuentas = []
        acreedor_seleccionado = None
        cuenta_actual = None
        propietarios = []
        historial = []

        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            termino = str(buscar_acreedor or "").strip()
            if termino:
                cur.execute(
                    """
                    SELECT identificacion, nombre, telefono, email
                    FROM contactos
                    WHERE tipo='Cliente'
                      AND (identificacion ILIKE %s OR nombre ILIKE %s)
                    ORDER BY nombre ASC
                    LIMIT 100
                    """,
                    (f"%{termino}%", f"%{termino}%"),
                )
                acreedores = [dict(r) for r in cur.fetchall()]

            if acreedor_id:
                cur.execute(
                    """
                    SELECT identificacion, nombre, telefono, email
                    FROM contactos
                    WHERE tipo='Cliente' AND identificacion=%s
                    LIMIT 1
                    """,
                    (str(acreedor_id).strip(),),
                )
                row = cur.fetchone()
                if row:
                    acreedor_seleccionado = dict(row)

            if acreedor_seleccionado:
                cur.execute(
                    """
                    SELECT
                        p.radicado_interno,
                        p.radicado_rama,
                        p.tipo_cartera,
                        p.naturaleza,
                        p.estado,
                        p.pretensiones,
                        p.inmueble_id,
                        p.id_demandado,
                        p.demandado,
                        i.conjunto_residencial,
                        i.torre_apto
                    FROM procesos p
                    LEFT JOIN inmuebles_ph i ON i.id=p.inmueble_id
                    WHERE %s = ANY(
                        string_to_array(
                            replace(COALESCE(p.id_cliente,''), ' ', ''),
                            '|'
                        )
                    )
                    ORDER BY p.radicado_interno DESC
                    """,
                    (acreedor_seleccionado["identificacion"],),
                )
                cuentas = [dict(r) for r in cur.fetchall()]

            if radicado_interno and acreedor_seleccionado:
                cuenta_actual = next(
                    (c for c in cuentas if str(c.get("radicado_interno")) == str(radicado_interno)),
                    None,
                )

            if cuenta_actual:
                inmueble_id = cuenta_actual.get("inmueble_id")

                if inmueble_id:
                    cur.execute(
                        """
                        SELECT c.identificacion, c.nombre, c.telefono, c.email, ip.es_principal
                        FROM inmueble_propietarios ip
                        JOIN contactos c ON c.id=ip.contacto_id
                        WHERE ip.inmueble_id=%s
                        ORDER BY ip.es_principal DESC, c.nombre ASC
                        """,
                        (int(inmueble_id),),
                    )
                    propietarios = [dict(r) for r in cur.fetchall()]

                if expedientes_service._table_exists(cur, "gestiones_crm"):
                    condiciones = ["radicado_interno=%s"]
                    params = [cuenta_actual["radicado_interno"]]
                    if inmueble_id:
                        condiciones.append("(radicado_interno IS NULL AND inmueble_id=%s)")
                        params.append(int(inmueble_id))

                    cur.execute(
                        f"""
                        SELECT id, tipo_contacto AS tipo,
                               identificacion_deudor AS deudor_nombre,
                               resumen, promesa_pago_fecha AS promesa,
                               usuario, fecha
                        FROM gestiones_crm
                        WHERE ({' OR '.join(condiciones)})
                          AND COALESCE(anulado,FALSE)=FALSE
                        ORDER BY fecha DESC
                        LIMIT 200
                        """,
                        params,
                    )
                    historial = []
                    for r in cur.fetchall():
                        d = dict(r)
                        d["tabla_origen"] = "gestiones_crm"
                        historial.append(d)

                if expedientes_service._table_exists(cur, "gestiones_cartera"):
                    ids_deudores = [
                        str(x).strip()
                        for x in str(cuenta_actual.get("id_demandado") or "").split("|")
                        if str(x).strip().isdigit()
                    ]
                    if ids_deudores:
                        cur.execute(
                            """
                            SELECT * FROM gestiones_cartera
                            WHERE REGEXP_REPLACE(
                                COALESCE(identificacion_deudor::text,''),
                                '[^0-9]', '', 'g'
                            ) = ANY(%s)
                            LIMIT 200
                            """,
                            (ids_deudores,),
                        )
                        for r in cur.fetchall():
                            d = dict(r)
                            if d.get("anulado") is True or str(d.get("estado", "")).upper() == "ANULADO":
                                continue
                            gestion_fecha = next(
                                (
                                    d.get(name)
                                    for name in (
                                        "fecha", "fecha_gestion", "created_at",
                                        "createdAt", "timestamp", "fecha_registro", "created"
                                    )
                                    if d.get(name) is not None
                                ),
                                None,
                            )
                            g_id = d.get("id") or d.get("id_gestion") or d.get("gestion_id")
                            historial.append({
                                "id": g_id,
                                "tipo": d.get("tipo_contacto") or "WhatsApp IA",
                                "deudor_nombre": d.get("identificacion_deudor") or "",
                                "resumen": d.get("resumen", ""),
                                "promesa": d.get("promesa_pago_fecha"),
                                "usuario": d.get("usuario", "Bot Claude"),
                                "fecha": gestion_fecha,
                                "tabla_origen": "gestiones_cartera",
                            })
                        historial.sort(key=lambda x: str(x.get("fecha") or ""), reverse=True)

        return render_template(
            "crm.html",
            {
                "request": request,
                "buscar_acreedor": buscar_acreedor or "",
                "acreedores": acreedores,
                "acreedor_seleccionado": acreedor_seleccionado,
                "cuentas": cuentas,
                "cuenta_actual": cuenta_actual,
                "inmueble_actual": cuenta_actual,
                "propietarios": propietarios,
                "historial": historial,
            },
        )
    finally:
        conn.release()


@app.post("/crm/guardar")
def crm_guardar(
    request: Request,
    radicado_interno: str | None = Form(None),
    inmueble_id: int | None = Form(None),
    tipo_contacto: str = Form(...),
    resumen: str = Form(...),
    promesa_pago_fecha: date | None = Form(None),
    identificacion_deudor: str | None = Form(None),
):
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO gestiones_crm
                        (
                            radicado_interno,
                            inmueble_id,
                            tipo_contacto,
                            resumen,
                            promesa_pago_fecha,
                            identificacion_deudor,
                            usuario
                        )
                    VALUES (%s, %s, %s, %s, %s, %s, 'ERP')
                    """,
                    (
                        str(radicado_interno or "").strip() or None,
                        inmueble_id,
                        tipo_contacto.strip(),
                        resumen.strip(),
                        promesa_pago_fecha,
                        identificacion_deudor,
                    ),
                )

        params = {"mensaje": "Gestión registrada"}
        if radicado_interno:
            params["radicado_interno"] = radicado_interno
        return _redirect("/crm", **params)
    except Exception as exc:
        print(f"[CRM] Error guardando gestion: {exc!r}", flush=True)
        params = {"error": "No fue posible guardar la gestión"}
        if radicado_interno:
            params["radicado_interno"] = radicado_interno
        return _redirect("/crm", **params)
    finally:
        conn.release()


@app.post("/crm/anular")
def crm_anular(
    gestion_id: int = Form(...),
    inmueble_id: str | None = Form(None),
    radicado_interno: str | None = Form(None),
    tabla_origen: str = Form("gestiones_crm"),
):
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                anulado = False

                if tabla_origen == "gestiones_crm":
                    cur.execute("SELECT id FROM gestiones_crm WHERE id=%s LIMIT 1", (gestion_id,))
                    if cur.fetchone():
                        cols = expedientes_service._cols(cur, "gestiones_crm")
                        if "estado" in cols:
                            cur.execute(
                                "UPDATE gestiones_crm SET anulado=TRUE, estado='ANULADO' WHERE id=%s",
                                (gestion_id,),
                            )
                        else:
                            cur.execute(
                                "UPDATE gestiones_crm SET anulado=TRUE WHERE id=%s",
                                (gestion_id,),
                            )
                        anulado = True

                if not anulado and expedientes_service._table_exists(cur, "gestiones_cartera"):
                    cols_c = expedientes_service._cols(cur, "gestiones_cartera")
                    id_col = "id" if "id" in cols_c else ("id_gestion" if "id_gestion" in cols_c else "gestion_id")
                    if id_col in cols_c:
                        cur.execute(
                            f"SELECT {id_col} FROM gestiones_cartera WHERE {id_col}=%s LIMIT 1",
                            (gestion_id,),
                        )
                        if cur.fetchone():
                            if "anulado" in cols_c:
                                cur.execute(
                                    f"UPDATE gestiones_cartera SET anulado=TRUE WHERE {id_col}=%s",
                                    (gestion_id,),
                                )
                            elif "estado" in cols_c:
                                cur.execute(
                                    f"UPDATE gestiones_cartera SET estado='ANULADO' WHERE {id_col}=%s",
                                    (gestion_id,),
                                )
                            else:
                                cur.execute(
                                    f"DELETE FROM gestiones_cartera WHERE {id_col}=%s",
                                    (gestion_id,),
                                )
                            anulado = True

                params = {}
                if radicado_interno:
                    params["radicado_interno"] = radicado_interno
                if not anulado:
                    return _redirect("/crm", error="Gestión no encontrada", **params)

        return _redirect("/crm", mensaje="Gestión anulada exitosamente", **params)
    except Exception as exc:
        print(f"[CRM] Error anulando gestion: {exc!r}", flush=True)
        params = {}
        if radicado_interno:
            params["radicado_interno"] = radicado_interno
        return _redirect("/crm", error="No fue posible anular la gestión", **params)
    finally:
        conn.release()


# ==============================================================================
# VENCIMIENTOS Y TÉRMINOS
# ==============================================================================
@app.get("/vencimientos")
def vencimientos(request: Request):
    conn = db.get_connection()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT radicado_interno FROM procesos ORDER BY radicado_interno DESC LIMIT 500")
            radicados = [r["radicado_interno"] for r in cur.fetchall()]
            cur.execute(
                "SELECT * FROM vencimientos WHERE completado=FALSE ORDER BY fecha_vencimiento ASC, id ASC"
            )
            pendientes = [dict(r) for r in cur.fetchall()]

            acuerdos = []
            if expedientes_service._table_exists(cur, "acuerdos_pago"):
                cur.execute("""
                    SELECT a.*, i.conjunto_residencial, i.torre_apto
                    FROM acuerdos_pago a
                    LEFT JOIN inmuebles_ph i ON a.inmueble_id = i.id
                    ORDER BY a.fecha_compromiso DESC LIMIT 200
                """)
                acuerdos = [dict(r) for r in cur.fetchall()]

            # Eventos unificados para el calendario
            eventos_calendario = []
            for v in pendientes:
                eventos_calendario.append({
                    "id": f"v-{v['id']}",
                    "tipo": "TERMINO_JUDICIAL",
                    "titulo": v.get("titulo") or "Término Judicial",
                    "fecha": str(v.get("fecha_vencimiento")),
                    "radicado": v.get("radicado_interno") or "",
                    "observaciones": v.get("observaciones") or "",
                    "completado": v.get("completado", False),
                })
            for a in acuerdos:
                val = float(a.get("valor_acordado") or 0)
                eventos_calendario.append({
                    "id": f"a-{a['id']}",
                    "tipo": "ACUERDO_PAGO",
                    "titulo": f"Pago: {a.get('nombre_deudor') or a.get('identificacion_deudor')} (${val:,.0f})",
                    "fecha": str(a.get("fecha_compromiso")),
                    "valor": val,
                    "estado": a.get("estado", "PENDIENTE"),
                    "deudor": a.get("nombre_deudor") or a.get("identificacion_deudor"),
                    "telefono": a.get("telefono") or "",
                    "inmueble": f"{a.get('conjunto_residencial') or ''} {a.get('torre_apto') or ''}".strip(),
                    "observaciones": a.get("observaciones") or "",
                })

        return render_template(
            "vencimientos.html",
            {
                "request": request,
                "radicados": radicados,
                "vencimientos": pendientes,
                "acuerdos": acuerdos,
                "json_eventos": json.dumps(eventos_calendario, ensure_ascii=False, default=str),
            },
        )
    finally:
        conn.release()


@app.post("/acuerdos/cumplir")
def cumplir_acuerdo(request: Request, acuerdo_id: int = Form(...)):
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute("UPDATE acuerdos_pago SET estado='CUMPLIDO', updated_at=CURRENT_TIMESTAMP WHERE id=%s", (acuerdo_id,))
                cur.execute("UPDATE vencimientos SET completado=TRUE WHERE tipo='ACUERDO_PAGO' AND observaciones ILIKE %s", (f"%Acuerdo #{acuerdo_id}%",))
        return _redirect("/vencimientos", mensaje="Acuerdo+marcado+como+cumplido")
    finally:
        conn.release()


@app.post("/acuerdos/incumplir")
def incumplir_acuerdo(request: Request, acuerdo_id: int = Form(...)):
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute("UPDATE acuerdos_pago SET estado='INCUMPLIDO', updated_at=CURRENT_TIMESTAMP WHERE id=%s", (acuerdo_id,))
        return _redirect("/vencimientos", mensaje="Acuerdo+marcado+como+incumplido")
    finally:
        conn.release()


@app.post("/acuerdos/guardar")
def guardar_acuerdo_manual(
    request: Request,
    identificacion_deudor: str = Form(...),
    fecha_compromiso: str = Form(...),
    valor_acordado: float = Form(0.0),
    nombre_deudor: str = Form(""),
    telefono: str = Form(""),
    inmueble_id: int | None = Form(None),
    observaciones: str = Form(""),
):
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute("""
                    INSERT INTO acuerdos_pago (
                        inmueble_id, identificacion_deudor, nombre_deudor, telefono,
                        valor_acordado, fecha_compromiso, estado, origen, observaciones
                    ) VALUES (%s, %s, %s, %s, %s, %s, 'PENDIENTE', 'ABOGADO_HUMANO', %s)
                    RETURNING id
                """, (
                    inmueble_id, identificacion_deudor.strip(), nombre_deudor.strip(), telefono.strip(),
                    valor_acordado, fecha_compromiso, observaciones.strip()
                ))
                acuerdo_id = cur.fetchone()[0]

                # Registrar en gestiones CRM
                cur.execute("""
                    INSERT INTO gestiones_crm (
                        inmueble_id, identificacion_deudor, tipo_contacto,
                        resumen, promesa_pago_fecha, usuario, estado
                    ) VALUES (%s, %s, 'Acuerdo Manual', %s, %s, 'Abogado ERP', 'ACTIVO')
                """, (
                    inmueble_id, identificacion_deudor.strip(),
                    f"🤝 [ACUERDO DE PAGO #{acuerdo_id}] Pactado por ${valor_acordado:,.0f} para el {fecha_compromiso}. {observaciones}",
                    fecha_compromiso
                ))

                # Registrar en vencimientos
                cur.execute("""
                    INSERT INTO vencimientos (
                        radicado_interno, titulo, fecha_vencimiento, observaciones,
                        completado, tipo, valor, inmueble_id
                    ) VALUES ('ACUERDO-PAGO', %s, %s, %s, FALSE, 'ACUERDO_PAGO', %s, %s)
                """, (
                    f"Pago acordado ({nombre_deudor or identificacion_deudor}) - ${valor_acordado:,.0f}",
                    fecha_compromiso,
                    f"Acuerdo #{acuerdo_id}. Tel: {telefono}. Obs: {observaciones}",
                    valor_acordado,
                    inmueble_id
                ))
        return _redirect("/vencimientos", mensaje="Acuerdo+registrado+exitosamente")
    finally:
        conn.release()


@app.post("/vencimientos/guardar")
def guardar_vencimiento(
    radicado_interno: str = Form(...),
    titulo: str = Form(...),
    fecha_vencimiento: date = Form(...),
    observaciones: str = Form(""),
):
    _ensure_crm_and_vencimientos_schema()
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute(
                    "INSERT INTO vencimientos (radicado_interno, titulo, fecha_vencimiento, observaciones) "
                    "VALUES (%s, %s, %s, %s)",
                    (radicado_interno.strip(), titulo.strip(), fecha_vencimiento, observaciones.strip()),
                )
        return RedirectResponse("/vencimientos", status_code=303)
    finally:
        conn.release()


@app.post("/vencimientos/completar")
def completar_vencimiento(vencimiento_id: int = Form(...)):
    _ensure_crm_and_vencimientos_schema()
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                cur.execute("UPDATE vencimientos SET completado=TRUE WHERE id=%s", (vencimiento_id,))
        return RedirectResponse("/vencimientos", status_code=303)
    finally:
        conn.release()


# ==============================================================================
# INFORMES Y DESCARGA EXCEL
# ==============================================================================
@app.get("/informes")
def informes(request: Request):
    conn = db.get_connection()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT COUNT(*) AS n FROM procesos")
            total = int(cur.fetchone()["n"])
            cur.execute("SELECT COUNT(*) AS n FROM contactos")
            total_contactos = int(cur.fetchone()["n"])
            cur.execute("SELECT COUNT(*) AS n FROM inmuebles_ph")
            total_inmuebles = int(cur.fetchone()["n"])
        return render_template("informes.html", {"request": request, "total_procesos": total, "total_contactos": total_contactos, "total_inmuebles": total_inmuebles})
    finally:
        conn.release()


@app.get("/descargar-excel")
def descargar_excel():
    conn = db.get_connection()
    try:
        output = exportaciones.generar_informe_ejecutivo_excel(conn)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=reporte_ejecutivo_expedientes.xlsx"},
        )
    finally:
        conn.release()


# ==============================================================================
# LIQUIDADOR DE EXPENSAS PH
# ==============================================================================
@app.get("/liquidador")
def vista_liquidador(
    request: Request,
    inmueble_id: str | None = None,
    tipo_tasa: str = "Máxima Legal",
    tasa_fija: float = 2.5,
    honorarios_pct: float = 23.8,
    gastos: float = 0.0,
):
    fecha_corte_obj = date.today()
    resultados = []
    resumen = {}
    lista_inmuebles = cargar_inmuebles_ph()
    inm_id_clean = None
    if inmueble_id:
        m = re.search(r"\d+", str(inmueble_id))
        if m:
            inm_id_clean = int(m.group(0))
    if inm_id_clean:
        try:
            resultados, resumen, _ = liquidador.motor_calculo_judicial(
                inm_id_clean, tipo_tasa, tasa_fija, honorarios_pct, gastos, fecha_corte_obj
            )
        except Exception as e:
            print(f"[LIQUIDADOR] Error en auto-cálculo: {e}", flush=True)
    return render_template("liquidador.html", {"request": request, "inmuebles": lista_inmuebles, "resultados": resultados, "resumen": resumen, "parametros": {"inmueble_id": inmueble_id, "tipo_tasa": tipo_tasa, "tasa_fija": tasa_fija, "honorarios_pct": honorarios_pct, "gastos": gastos, "fecha_corte": fecha_corte_obj.strftime("%Y-%m-%d")}})


@app.post("/liquidador")
def calcular_liquidador(
    request: Request,
    inmueble_id: int = Form(...),
    tipo_tasa: str = Form(...),
    tasa_fija: float = Form(2.5),
    honorarios_pct: float = Form(23.8),
    gastos: float = Form(0.0),
    fecha_corte: date = Form(...),
):
    resultados, resumen, _ = liquidador.motor_calculo_judicial(
        inmueble_id, tipo_tasa, tasa_fija, honorarios_pct, gastos, fecha_corte
    )
    if not resultados:
        return render_template("liquidador.html", {"request": request, "inmuebles": cargar_inmuebles_ph(), "error": "No hay deudas.", "resultados": None})
    return render_template("liquidador.html", {"request": request, "inmuebles": cargar_inmuebles_ph(), "resultados": resultados, "resumen": resumen, "parametros": {"inmueble_id": inmueble_id, "tipo_tasa": tipo_tasa, "tasa_fija": tasa_fija, "honorarios_pct": honorarios_pct, "gastos": gastos, "fecha_corte": fecha_corte.strftime("%Y-%m-%d")}})


@app.post("/liquidador/actualizar")
async def actualizar_cuotas(request: Request):
    form_data = await request.form()
    inmueble_id = int(form_data.get("inmueble_id"))
    conn = db.get_connection()
    try:
        with conn:
            with conn.cursor() as cur:
                for key, value in form_data.items():
                    if key.startswith(("ord_", "ext_", "gas_", "abo_")):
                        prefijo, y, m = key.split("_")
                        concepto = {
                            "ord": "Expensa Ordinaria",
                            "ext": "Cuota Extraordinaria",
                            "gas": "Gastos",
                            "abo": "Abono",
                        }[prefijo]
                        valor = float(value) if value else 0.0
                        cur.execute("""
                            UPDATE expensas_ph
                            SET valor_capital = %s
                            WHERE inmueble_id = %s AND concepto = %s
                              AND periodo_anio = %s AND periodo_mes = %s
                        """, (valor, inmueble_id, concepto, int(y), int(m)))
                        if cur.rowcount == 0 and valor > 0:
                            f_vencimiento = f"{y}-{int(m):02d}-01"
                            cur.execute("""
                                INSERT INTO expensas_ph
                                    (inmueble_id, concepto, periodo_mes, periodo_anio,
                                     valor_capital, fecha_vencimiento, estado)
                                VALUES (%s, %s, %s, %s, %s, %s, 'Aplicado')
                            """, (inmueble_id, concepto, int(m), int(y), valor, f_vencimiento))
    except Exception as e:
        print(f"[LIQUIDADOR] Error actualizando cuotas: {e}", flush=True)
    finally:
        conn.release()
    return RedirectResponse(url="/liquidador", status_code=307)


@app.post("/liquidador/exportar/pdf")
async def exportar_pdf(
    request: Request,
    inmueble_id: int = Form(...),
    tipo_tasa: str = Form(...),
    tasa_fija: float = Form(2.5),
    honorarios_pct: float = Form(23.8),
    gastos: float = Form(0.0),
    fecha_corte=None,
):
    if isinstance(fecha_corte, str):
        try:
            fecha_corte = datetime.strptime(fecha_corte, "%Y-%m-%d").date()
        except ValueError:
            fecha_corte = date.today()
    if fecha_corte is None:
        fecha_corte = date.today()

    resultados, resumen, inm_info = liquidador.motor_calculo_judicial(
        inmueble_id, tipo_tasa, tasa_fija, honorarios_pct, gastos, fecha_corte
    )
    path = exportaciones.generar_pdf_liquidacion(inmueble_id, fecha_corte, resultados, resumen, inm_info)
    return FileResponse(
        path,
        media_type="application/pdf",
        filename=f"Liquidacion_{inmueble_id}_{fecha_corte.isoformat()}.pdf",
    )


@app.post("/liquidador/exportar/excel")
async def exportar_excel(
    request: Request,
    inmueble_id: int = Form(...),
    tipo_tasa: str = Form(...),
    tasa_fija: float = Form(2.5),
    honorarios_pct: float = Form(23.8),
    gastos: float = Form(0.0),
    fecha_corte: date = Form(...),
):
    resultados, resumen, inm_info = liquidador.motor_calculo_judicial(
        inmueble_id, tipo_tasa, tasa_fija, honorarios_pct, gastos, fecha_corte
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Liquidación"
    ws['A1'] = "LIQUIDACIÓN DE CRÉDITO - PROPIEDAD HORIZONTAL"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"FECHA DE GENERACIÓN: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws['A3'] = f"DEMANDANTE / CONJUNTO: {inm_info[0]} - {inm_info[1]}"
    ws['A4'] = f"DEUDOR: {inm_info[2]} (CC/NIT: {inm_info[3]})"
    ws['A5'] = f"FECHA DE CORTE: {fecha_corte.strftime('%Y-%m-%d')}"
    headers = ['Período', 'Ordinaria', 'Extraord.', 'Gastos', 'Abonos', 'Cap. Liquidable', 'Días', 'Tasa E.A.', 'Tasa Mensual', 'Interés Mes', 'Int. Acumulado', 'Saldo Final']
    ws.append([])
    ws.append(headers)
    for cell in ws[7]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    row_idx = 8
    for r in resultados:
        ws.cell(row=row_idx, column=1, value=r['desde'][:7])
        ws.cell(row=row_idx, column=2, value=r['ordinarias'])
        ws.cell(row=row_idx, column=3, value=r['extraordinarias'])
        ws.cell(row=row_idx, column=4, value=r['gastos'])
        ws.cell(row=row_idx, column=5, value=r['abonos'])
        ws.cell(row=row_idx, column=7, value=r['dias'])
        ws.cell(row=row_idx, column=8, value=float(r['tasa_ea'].replace('%', '')) / 100.0)
        ws.cell(row=row_idx, column=9, value=f"=(1+H{row_idx})^(1/12)-1")
        if row_idx == 8:
            ws.cell(row=row_idx, column=10, value=f"=ROUND((B{row_idx}+C{row_idx}+D{row_idx})*I{row_idx}*(G{row_idx}/30), 2)")
            ws.cell(row=row_idx, column=11, value=f"=MAX(0, J{row_idx}-E{row_idx})")
            ws.cell(row=row_idx, column=6, value=f"=(B{row_idx}+C{row_idx}+D{row_idx})-MAX(0, E{row_idx}-J{row_idx})")
            ws.cell(row=row_idx, column=12, value=f"=F{row_idx}+K{row_idx}")
        else:
            ws.cell(row=row_idx, column=10, value=f"=ROUND((F{row_idx-1}+B{row_idx}+C{row_idx}+D{row_idx})*I{row_idx}*(G{row_idx}/30), 2)")
            ws.cell(row=row_idx, column=11, value=f"=MAX(0, K{row_idx-1}+J{row_idx}-E{row_idx})")
            ws.cell(row=row_idx, column=6, value=f"=(F{row_idx-1}+B{row_idx}+C{row_idx}+D{row_idx})-MAX(0, E{row_idx}-(K{row_idx-1}+J{row_idx}))")
            ws.cell(row=row_idx, column=12, value=f"=F{row_idx}+K{row_idx}")
        for col in [2,3,4,5,6,10,11,12]:
            ws.cell(row=row_idx, column=col).number_format = '"$"#,##0'
        for col in [8,9]:
            ws.cell(row=row_idx, column=col).number_format = '0.0000%'
        row_idx += 1
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = max_length + 2
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    safe_name = str(inm_info[2] or "Deudor").replace(" ", "_")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Liquidacion_{safe_name}.xlsx"},
    )


@app.get("/liquidador/plantilla")
def descargar_plantilla_liquidador():
    datos_ejemplo = {
        "Desde": ["01/01/2025", "01/02/2025"],
        "Hasta": ["31/01/2025", "28/02/2025"],
        "ORDINARIAS": [250000, 250000],
        "EXTRAORDINARIAS": [0, 50000],
    }
    df = pd.DataFrame(datos_ejemplo)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Cuotas")
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=formato_cuotas_ph.xlsx"},
    )


@app.post("/liquidador/carga_masiva")
async def carga_masiva_excel(
    request: Request,
    inmueble_id: int = Form(...),
    tipo_tasa: str = Form(...),
    tasa_fija: float = Form(2.5),
    honorarios_pct: float = Form(23.8),
    gastos: float = Form(0.0),
    fecha_corte: date = Form(...),
    archivo: UploadFile = File(None),
):
    if archivo and archivo.filename:
        try:
            contents = await archivo.read()
            df = pd.read_excel(io.BytesIO(contents))
            df.columns = df.columns.str.strip().str.lower().str.replace('ñ', 'n')
            if 'desde' in df.columns:
                meses_es = {
                    'ene': '01', 'feb': '02', 'mar': '03', 'abr': '04',
                    'may': '05', 'jun': '06', 'jul': '07', 'ago': '08',
                    'sep': '09', 'oct': '10', 'nov': '11', 'dic': '12',
                }

                def parse_spanish_date(d):
                    if pd.isna(d):
                        return None
                    if isinstance(d, (datetime, date)):
                        return d
                    d_str = str(d).lower().strip()
                    for es, num in meses_es.items():
                        d_str = d_str.replace(es, f"-{num}-")
                    try:
                        return pd.to_datetime(d_str, dayfirst=True, errors='coerce')
                    except Exception:
                        return None

                df['desde'] = df['desde'].apply(parse_spanish_date)
                cuotas_procesadas = 0
                conn = db.get_connection()
                try:
                    with conn:
                        with conn.cursor() as cur:
                            for index, row in df.iterrows():
                                fecha = row['desde']
                                if pd.isna(fecha):
                                    continue
                                y, m = fecha.year, fecha.month
                                ord_val = row.get('ordinarias', row.get('ordinaria', 0))
                                ext_val = row.get('extraordinarias', row.get('extraordinaria', 0))
                                conceptos = {'Expensa Ordinaria': ord_val, 'Cuota Extraordinaria': ext_val}
                                for concepto, valor in conceptos.items():
                                    try:
                                        if pd.isna(valor) or str(valor).strip() == '':
                                            valor_limpio = 0.0
                                        elif isinstance(valor, (int, float)):
                                            valor_limpio = float(valor)
                                        else:
                                            v_str = str(valor).replace('$', '').strip()
                                            if '.' in v_str and ',' in v_str:
                                                v_str = v_str.replace('.', '').replace(',', '.')
                                            elif ',' in v_str:
                                                v_str = v_str.replace(',', '.')
                                            valor_limpio = float(v_str)
                                    except Exception:
                                        valor_limpio = 0.0
                                    if valor_limpio > 0:
                                        cur.execute("""
                                            UPDATE expensas_ph
                                            SET valor_capital = %s
                                            WHERE inmueble_id = %s AND concepto = %s
                                              AND periodo_anio = %s AND periodo_mes = %s
                                        """, (valor_limpio, inmueble_id, concepto, y, m))
                                        if cur.rowcount == 0:
                                            f_vencimiento = f"{y}-{m:02d}-01"
                                            cur.execute("""
                                                INSERT INTO expensas_ph
                                                    (inmueble_id, concepto, periodo_mes, periodo_anio,
                                                     valor_capital, fecha_vencimiento, estado)
                                                VALUES (%s, %s, %s, %s, %s, %s, 'En Mora')
                                            """, (inmueble_id, concepto, m, y, valor_limpio, f_vencimiento))
                                            cuotas_procesadas += 1
                finally:
                    conn.release()
        except Exception as e:
            print(f"[LIQUIDADOR] Error procesando carga masiva: {e}", flush=True)

    resultados, resumen, _ = liquidador.motor_calculo_judicial(
        inmueble_id, tipo_tasa, tasa_fija, honorarios_pct, gastos, fecha_corte
    )
    lista_inmuebles = cargar_inmuebles_ph()
    return render_template("liquidador.html", {"request": request, "inmuebles": lista_inmuebles, "resultados": resultados, "resumen": resumen, "parametros": {"inmueble_id": inmueble_id, "tipo_tasa": tipo_tasa, "tasa_fija": tasa_fija, "honorarios_pct": honorarios_pct, "gastos": gastos, "fecha_corte": fecha_corte.strftime("%Y-%m-%d")}})
