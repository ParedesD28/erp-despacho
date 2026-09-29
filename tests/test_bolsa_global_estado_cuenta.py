"""
Integración Bolsa Global ↔ lote estado de cuenta (servicio + caché).

No usa TestClient/httpx (CI del ERP no lo instala de forma fiable).
"""
from __future__ import annotations

import io
import json
import unittest

from bolsa_global_estado_cuenta_service import (
    calcular_bolsa_lote,
    calcular_bolsa_por_cuenta,
    generar_excel_bolsa_global,
    resultado_a_json_bytes,
)
from estado_cuenta_pdf_service import (
    cuentas_desde_cache,
    excel_desde_cache,
    generar_excel_lote,
    procesar_lote_estados_cuenta,
)


def _fila(fecha, concepto, valor=0, abono=0, **extra):
    base = {
        "Concepto": concepto,
        "Tipo Documento": "FAC",
        "Número": "1",
        "Fecha": fecha,
        "Valor": valor,
        "Abono": abono,
        "Saldo": 0,
        "Titular": extra.get("titular", "DEUDOR A"),
        "Bloque": extra.get("bloque", "1"),
        "Apartamento": extra.get("apartamento", "101"),
        "Codigo Cuenta": extra.get("codigo", "C1"),
        "Archivo": extra.get("archivo", "a.pdf"),
    }
    return base


def _cuenta(archivo, titular, rows, error=None):
    return {
        "archivo": archivo,
        "titular": titular,
        "bloque": "1",
        "apartamento": "101",
        "codigo_cuenta": "C1",
        "conjunto": "CONJUNTO",
        "movimientos_extraidos": len(rows),
        "rows": rows,
        "error": error,
    }


class BolsaIntegracionServicioTests(unittest.TestCase):
    def test_por_cuenta_aplica_bolsa_global(self):
        rows = [
            _fila("2024.01.01", "CUOTA ADMON", 100, 0),
            _fila("2024.01.01", "INTERES MORA", 20, 0),
            _fila("2024.01.15", "ABONO", 0, 10),
            _fila("2024.02.01", "CUOTA ADMON", 100, 0),
        ]
        out = calcular_bolsa_por_cuenta(_cuenta("a.pdf", "DEUDOR A", rows))
        self.assertEqual(out["archivo"], "a.pdf")
        self.assertEqual(out["titular"], "DEUDOR A")
        self.assertFalse(out["omitido"])
        self.assertIn("fecha_inicio_mora", out)
        self.assertIn("capital_limpio_a_demandar", out)
        self.assertIn("bolsa_global_inicial", out)
        self.assertIn("bolsa_remanente_final", out)
        self.assertIn("errores_procesamiento", out)
        self.assertEqual(out["bolsa_global_inicial"], 10.0)
        self.assertTrue(out["total_capital_demandado"] > 0)

    def test_lote_no_mezcla_deudores(self):
        c1 = _cuenta(
            "a.pdf",
            "ANA",
            [
                _fila("2024.01.01", "CUOTA ADMON", 100, 0, titular="ANA", archivo="a.pdf"),
                _fila("2024.01.01", "INTERES", 50, 0, titular="ANA", archivo="a.pdf"),
            ],
        )
        c2 = _cuenta(
            "b.pdf",
            "BETO",
            [
                _fila("2024.03.01", "CUOTA ADMON", 200, 0, titular="BETO", archivo="b.pdf"),
                _fila("2024.03.10", "PAGO", 0, 200, titular="BETO", archivo="b.pdf"),
            ],
        )
        lote = calcular_bolsa_lote([c1, c2])
        self.assertEqual(lote["cuentas_evaluadas"], 2)
        self.assertEqual(len(lote["resultados"]), 2)
        por_archivo = {r["archivo"]: r for r in lote["resultados"]}
        self.assertEqual(por_archivo["a.pdf"]["titular"], "ANA")
        self.assertEqual(por_archivo["b.pdf"]["titular"], "BETO")
        # BETO pagó todo el capital del mes → sin capital limpio típico / bolsa consumida
        self.assertEqual(por_archivo["a.pdf"]["bolsa_global_inicial"], 0.0)
        self.assertEqual(por_archivo["b.pdf"]["bolsa_global_inicial"], 200.0)
        # Totales de lote = suma de totales individuales
        self.assertAlmostEqual(
            lote["total_capital_demandado_lote"],
            por_archivo["a.pdf"]["total_capital_demandado"]
            + por_archivo["b.pdf"]["total_capital_demandado"],
            places=2,
        )

    def test_cuenta_con_error_parseo_queda_omitida(self):
        out = calcular_bolsa_por_cuenta(
            _cuenta("malo.pdf", "", [], error="PDF corrupto")
        )
        self.assertTrue(out["omitido"])
        self.assertEqual(out["fecha_inicio_mora"], "Sin deuda")
        self.assertTrue(any("PDF no parseado" in e for e in out["errores_procesamiento"]))

    def test_excel_y_json_export_separados(self):
        lote = calcular_bolsa_lote(
            [
                _cuenta(
                    "a.pdf",
                    "ANA",
                    [
                        _fila("2024.01.01", "CUOTA ADMON", 100, 0),
                        _fila("2024.01.01", "INTERES MORA", 30, 0),
                    ],
                )
            ]
        )
        xlsx = generar_excel_bolsa_global(lote)
        self.assertTrue(xlsx.getvalue().startswith(b"PK"))
        raw = resultado_a_json_bytes(lote)
        data = json.loads(raw.getvalue().decode("utf-8"))
        self.assertEqual(data["cuentas_evaluadas"], 1)
        self.assertIn("resultados", data)


class BolsaCacheYTranscripcionTests(unittest.TestCase):
    def test_cuentas_desde_cache_tras_analizar(self):
        # PDF inválido: queda en lote con error, pero la caché guarda la cuenta.
        resultado = procesar_lote_estados_cuenta(
            [("malo.pdf", b"no-es-pdf")],
            incluir_excel=False,
        )
        self.assertTrue(resultado["cache_id"])
        cuentas = cuentas_desde_cache(resultado["cache_id"])
        self.assertIsNotNone(cuentas)
        self.assertEqual(len(cuentas), 1)
        self.assertEqual(cuentas[0]["archivo"], "malo.pdf")
        self.assertIn("rows", cuentas[0])

        bolsa = calcular_bolsa_lote(cuentas)
        self.assertEqual(bolsa["cuentas_evaluadas"], 1)
        self.assertTrue(bolsa["resultados"][0]["omitido"])

    def test_excel_transcripcion_no_incluye_hojas_bolsa(self):
        """PR #20: generar_excel_lote sigue siendo solo transcripción."""
        cuentas = [
            {
                "archivo": "a.pdf",
                "titular": "ANA",
                "bloque": "1",
                "apartamento": "101",
                "codigo_cuenta": "C1",
                "conjunto": "X",
                "fechas_detectadas": 1,
                "movimientos_extraidos": 1,
                "bloques_omitidos": 0,
                "inconsistencias_saldo": 0,
                "saldo_final": 100,
                "error": None,
                "rows": [_fila("2024.01.01", "CUOTA ADMON", 100, 0)],
            }
        ]
        excel = generar_excel_lote(cuentas)
        from openpyxl import load_workbook

        wb = load_workbook(excel)
        nombres = set(wb.sheetnames)
        self.assertIn("Resumen", nombres)
        self.assertIn("Movimientos", nombres)
        self.assertNotIn("Resumen Bolsa", nombres)
        self.assertNotIn("Capital limpio", nombres)

        bolsa_xlsx = generar_excel_bolsa_global(calcular_bolsa_lote(cuentas))
        wb2 = load_workbook(bolsa_xlsx)
        self.assertIn("Resumen Bolsa", wb2.sheetnames)
        self.assertIn("Capital limpio", wb2.sheetnames)

    def test_tras_armar_excel_cache_conserva_cuentas(self):
        """Excel cachea bytes pero deja rows para Bolsa/certificados (anti re-parseo)."""
        resultado = procesar_lote_estados_cuenta(
            [("malo.pdf", b"no-es-pdf")],
            incluir_excel=False,
        )
        cache_id = resultado["cache_id"]
        self.assertIsNotNone(cuentas_desde_cache(cache_id))
        libro = excel_desde_cache(cache_id)
        self.assertIsNotNone(libro)
        self.assertTrue(libro.getvalue().startswith(b"PK"))
        # Rows siguen disponibles tras Excel de transcripción.
        self.assertIsNotNone(cuentas_desde_cache(cache_id))
        # Segunda descarga Excel reusa bytes cacheados.
        otra = excel_desde_cache(cache_id)
        self.assertTrue(otra.getvalue().startswith(b"PK"))


class BolsaPermisosRutaTests(unittest.TestCase):
    def test_post_bolsa_exige_accion_editar(self):
        import permisos

        consulta = permisos.permisos_de_perfil(permisos.PERFIL_CONSULTA)
        self.assertTrue(
            permisos.denegar_acceso(consulta, "POST", "/herramientas/estado-cuenta/bolsa-global")
        )
        self.assertTrue(
            permisos.denegar_acceso(
                consulta, "POST", "/herramientas/estado-cuenta/bolsa-global/excel"
            )
        )
        abogado = permisos.permisos_de_perfil(permisos.PERFIL_ABOGADO)
        self.assertFalse(
            permisos.denegar_acceso(abogado, "POST", "/herramientas/estado-cuenta/bolsa-global")
        )

    def test_get_herramientas_exige_nav_informes(self):
        import permisos

        consulta = permisos.permisos_de_perfil(permisos.PERFIL_CONSULTA)
        self.assertFalse(
            permisos.denegar_acceso(consulta, "GET", "/herramientas/estado-cuenta")
        )
        # Perfil sin nav.informes (solo cobro parcial) no debería ver la herramienta.
        sin_informes = frozenset({permisos.NAV_COBRO_CRM, permisos.ACCION_EDITAR})
        self.assertTrue(
            permisos.denegar_acceso(sin_informes, "GET", "/herramientas/estado-cuenta")
        )


class BolsaResultadoVacioTests(unittest.TestCase):
    def test_resultado_vacio_contrato(self):
        from bolsa_global_mora import resultado_vacio

        r = resultado_vacio(errores=["x"])
        self.assertEqual(r["fecha_inicio_mora"], "Sin deuda")
        self.assertEqual(r["capital_limpio_a_demandar"], [])
        self.assertEqual(r["errores_procesamiento"], ["x"])


if __name__ == "__main__":
    unittest.main()
