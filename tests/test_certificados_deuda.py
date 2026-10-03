"""Tests del módulo Certificados de Deuda (agrupación, matching Neon, Word)."""
from __future__ import annotations

import sys
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import certificados_deuda_repository as repo
import certificados_deuda_service as svc
import permisos


class AgruparCapitalTests(unittest.TestCase):
    def test_agrupa_ordinarias_y_extras_con_saldo_acumulado(self):
        items = [
            {
                "fecha": "2024.02.10",
                "concepto": "CUOTA ADMINISTRACION",
                "valor_a_demandar": 100.0,
                "clase_concepto": "capital",
            },
            {
                "fecha": "2024.02.15",
                "concepto": "CUOTA EXTRAORDINARIA ASCENSOR",
                "valor_a_demandar": 50.0,
                "clase_concepto": "capital",
            },
            {
                "fecha": "2024.03.01",
                "concepto": "CUOTA ADMON",
                "valor_a_demandar": 200.0,
                "clase_concepto": "capital",
            },
        ]
        filas = svc.agrupar_capital_limpio(items)
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[0].mes, "FEBRERO")
        self.assertEqual(filas[0].anio, 2024)
        self.assertEqual(filas[0].periodo, "FEBRERO 2024")
        self.assertEqual(filas[0].cuotas_ordinarias, 100.0)
        self.assertEqual(filas[0].cuotas_extraordinarias, 50.0)
        self.assertEqual(filas[0].total_mes, 150.0)
        self.assertEqual(filas[0].saldo, 150.0)
        self.assertEqual(filas[0].vencimiento, "29-feb-24")  # 2024 bisiesto
        self.assertEqual(filas[1].mes, "MARZO")
        self.assertEqual(filas[1].cuotas_ordinarias, 200.0)
        self.assertEqual(filas[1].cuotas_extraordinarias, 0.0)
        self.assertEqual(filas[1].saldo, 350.0)
        self.assertEqual(filas[1].vencimiento, "31-mar-24")

    def test_gasto_va_a_extraordinarias(self):
        items = [
            {
                "fecha": "2024.04.01",
                "concepto": "HONORARIOS PREJURIDICO",
                "valor_a_demandar": 80.0,
                "clase_concepto": "gasto",
            }
        ]
        filas = svc.agrupar_capital_limpio(items)
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0].cuotas_ordinarias, 0.0)
        self.assertEqual(filas[0].cuotas_extraordinarias, 80.0)

    def test_mes_corte_como_fallback_de_fecha(self):
        items = [
            {
                "mes_corte": "2024.05",
                "concepto": "CUOTA ADMINISTRATIVA",
                "valor_a_demandar": 10.0,
            }
        ]
        filas = svc.agrupar_capital_limpio(items)
        self.assertEqual(filas[0].periodo, "MAYO 2024")
        self.assertEqual(filas[0].vencimiento, "31-may-24")

    def test_vencimiento_fin_mes_ene_feb_bisiesto_y_override(self):
        ene = svc.agrupar_capital_limpio(
            [{"fecha": "2024.01.10", "concepto": "CUOTA ADMINISTRACION", "valor_a_demandar": 1.0}]
        )
        feb_bisiesto = svc.agrupar_capital_limpio(
            [{"fecha": "2024.02.10", "concepto": "CUOTA ADMINISTRACION", "valor_a_demandar": 1.0}]
        )
        feb_comun = svc.agrupar_capital_limpio(
            [{"fecha": "2023.02.10", "concepto": "CUOTA ADMINISTRACION", "valor_a_demandar": 1.0}]
        )
        override = svc.agrupar_capital_limpio(
            [{"fecha": "2024.01.10", "concepto": "CUOTA ADMINISTRACION", "valor_a_demandar": 1.0}],
            dia_vencimiento=5,
        )
        self.assertEqual(ene[0].vencimiento, "31-ene-24")
        self.assertEqual(feb_bisiesto[0].vencimiento, "29-feb-24")
        self.assertEqual(feb_comun[0].vencimiento, "28-feb-23")
        self.assertEqual(override[0].vencimiento, "5-ene-24")


class MatchingUnidadTests(unittest.TestCase):
    """Regresión del 404: PDF 'TORRE X APTO Y' vs Neon 'X-Y' / 'TX-Y'."""

    def test_clave_une_torre_apto_texto_y_bloque_apartamento(self):
        self.assertEqual(repo._clave_unidad("TORRE 1 APTO 201"), "1-201")
        self.assertEqual(repo._clave_unidad("1-201"), "1-201")
        self.assertEqual(
            repo._clave_unidad("", bloque="1", apartamento="201"),
            "1-201",
        )
        self.assertEqual(
            repo._clave_unidad("TORRE 32 APTO N0.502"),
            "32-502",
        )
        self.assertEqual(repo._clave_unidad("BLOQUE 9 APTO 401"), "9-401")

    def test_claves_cruzan_formatos_tipicos_colon_vs_neon(self):
        pdf = repo._claves_unidad(
            "TORRE 1 APTO 201",
            bloque="1",
            apartamento="201",
        )
        neon_simple = repo._claves_unidad("1-201")
        neon_t = repo._claves_unidad("T1-201")
        neon_largo = repo._claves_unidad("TORRE 1 APTO 201")
        self.assertTrue(pdf & neon_simple)
        self.assertTrue(pdf & neon_t)
        self.assertTrue(pdf & neon_largo)
        self.assertTrue(repo._unidad_coincide("1-201", pdf))
        self.assertTrue(repo._unidad_coincide("T1-201", pdf))
        self.assertTrue(repo._unidad_coincide("TORRE 1 APTO 201", pdf))
        self.assertFalse(repo._unidad_coincide("1-202", pdf))

    def test_ceros_izquierda_02_042_equivale_2_42(self):
        """Regresión Mirador: PDF TORRE 02 APTO 042 vs Neon 2-42 / 02-42."""
        pdf = repo._claves_unidad(
            "TORRE 02 APTO 042",
            bloque="02",
            apartamento="042",
        )
        self.assertIn("02-042", pdf)
        self.assertIn("2-42", pdf)
        self.assertIn("02-42", pdf)
        self.assertIn("2-042", pdf)
        self.assertTrue(pdf & repo._claves_unidad("2-42"))
        self.assertTrue(pdf & repo._claves_unidad("02-42"))
        self.assertTrue(repo._unidad_coincide("2-42", pdf))
        self.assertTrue(repo._unidad_coincide("02-42", pdf))
        self.assertTrue(repo._unidad_coincide("02-042", pdf))
        self.assertFalse(repo._unidad_coincide("2-43", pdf))
        self.assertEqual(repo._tupla_numerica_unidad("02-042"), (2, 42))
        self.assertEqual(repo._tupla_numerica_unidad("2-42"), (2, 42))

    def test_unidad_no_match_por_sufijo_prefijo_ni_compacto(self):
        """Regla general: unidades distintas no matchean por sufijo/prefijo/pegado.

        #52 quitó substring en `_unidad_coincide`, pero `_variantes_clave_unidad`
        aún pegaba torre+apto (`1-431`→`1431`). Este test cubre ambos vectores.
        """
        casos = (
            # Apto plano sufijo/prefijo
            ("1431", "431"),
            ("431", "1431"),
            ("1124", "124"),
            ("124", "1124"),
            ("1101", "101"),
            ("101", "1101"),
            ("12", "2"),
            ("2", "12"),
            ("1001", "001"),
            ("001", "1001"),
            ("1001", "1"),
            ("201", "01"),
            ("01", "201"),
            ("201", "1"),
            # Compacto torre+apto ≢ apto plano (causa raíz 1431↔431)
            ("1431", "1-431"),
            ("1-431", "1431"),
            ("1124", "1-124"),
            ("1-124", "1124"),
            ("1101", "1-101"),
            ("1-101", "1101"),
            ("242", "2-42"),
            ("2-42", "242"),
            ("101", "1-01"),
            ("1-01", "101"),
            # Compuestos misma torre, apto sufijo
            ("1-1431", "1-431"),
            ("1-431", "1-1431"),
            ("1-124", "1-24"),
            ("1-24", "1-124"),
            ("1-1101", "1-101"),
            ("1-101", "1-1101"),
            ("TORRE 1 APTO 1124", "1-124"),
            ("TORRE 1 APTO 124", "1-1124"),
            ("TORRE 1 APTO 1431", "1-431"),
            ("TORRE 1 APTO 431", "1-1431"),
            # Torre letra T-A / TA
            ("TA-1431", "TA-431"),
            ("TA-431", "TA-1431"),
            ("T-A-1431", "T-A-431"),
            ("TA-1431", "431"),
            ("431", "TA-1431"),
            ("TA-1431", "1-431"),
            ("TORRE TA APTO 1431", "TA-431"),
            ("TORRE TA APTO 431", "TA-1431"),
        )
        for neon, busqueda in casos:
            with self.subTest(neon=neon, busqueda=busqueda):
                claves = repo._claves_unidad(busqueda)
                self.assertFalse(
                    repo._unidad_coincide(neon, claves),
                    msg=f"falso positivo: {neon!r} vs {busqueda!r}",
                )
                self.assertFalse(
                    repo._unidad_coincide(busqueda, repo._claves_unidad(neon)),
                    msg=f"falso positivo inverso: {busqueda!r} vs {neon!r}",
                )

    def test_unidad_match_igualdad_legitima(self):
        """Positivos: alias de formato, ceros a la izquierda, torre letra."""
        casos = (
            ("1-201", "TORRE 1 APTO 201"),
            ("1-201", "T1-201"),
            ("T1-201", "1-201"),
            ("02-042", "2-42"),
            ("2-42", "02-042"),
            ("02-042", "TORRE 02 APTO 042"),
            ("0124", "124"),
            ("124", "0124"),
            ("TA-1431", "T-A-1431"),
            ("T-A-1431", "TA-1431"),
            ("TA-01431", "TA-1431"),
            ("TORRE TA APTO 1431", "TA-1431"),
            ("TA-431", "T-A-431"),
            ("9-401", "BLOQUE 9 APTO 401"),
        )
        for neon, busqueda in casos:
            with self.subTest(neon=neon, busqueda=busqueda):
                self.assertTrue(
                    repo._unidad_coincide(neon, repo._claves_unidad(busqueda)),
                    msg=f"debía coincidir: {neon!r} vs {busqueda!r}",
                )

    def test_variantes_no_pegan_torre_apto_en_digitos(self):
        """1-431 no genera clave 1431; 2-42 no genera 242."""
        self.assertNotIn("1431", repo._variantes_clave_unidad("1-431"))
        self.assertNotIn("1124", repo._variantes_clave_unidad("1-124"))
        self.assertNotIn("242", repo._variantes_clave_unidad("2-42"))
        self.assertIn("124", repo._variantes_clave_unidad("0124"))
        self.assertIn("2-42", repo._variantes_clave_unidad("02-042"))

    def test_clave_preserva_torre_solo_letras(self):
        self.assertEqual(repo._clave_unidad("T-A-1431"), "TA-1431")
        self.assertEqual(
            repo._clave_unidad("", bloque="T-A", apartamento="1431"),
            "TA-1431",
        )
        self.assertEqual(
            repo.clave_canonica_unidad("", bloque="TA", apartamento="0431"),
            "TA-431",
        )
        self.assertEqual(
            repo._claves_unidad("", bloque="T-A", apartamento="431"),
            {"TA-431"},
        )

    def test_codigo_cuenta_compatible_anti_sufijo_colon(self):
        """Identificación COLON 1431: OK con 1-431/1431; no con apto plano 431."""
        self.assertTrue(repo._compatible_con_codigo_cuenta("1-431", "1431"))
        self.assertTrue(repo._compatible_con_codigo_cuenta("T1-431", "1431"))
        self.assertTrue(repo._compatible_con_codigo_cuenta("1431", "1431"))
        self.assertTrue(repo._compatible_con_codigo_cuenta("01431", "1431"))
        self.assertFalse(repo._compatible_con_codigo_cuenta("431", "1431"))
        self.assertFalse(repo._compatible_con_codigo_cuenta("0431", "1431"))
        self.assertTrue(repo._compatible_con_codigo_cuenta("1-124", "1124"))
        self.assertFalse(repo._compatible_con_codigo_cuenta("124", "1124"))
        # Torre letra + apto: pegado no aplica; exacto por claves de código
        self.assertFalse(repo._compatible_con_codigo_cuenta("TA-431", "1431"))
        self.assertTrue(repo._compatible_con_codigo_cuenta("TA-1431", "TA1431"))

    def test_elegir_prioriza_codigo_cuenta_sobre_sufijo_apto(self):
        """PDF COLON: codigo=1431 bloque=1 apto=431 no debe caer en Neon 431."""
        filas = [
            {
                "inmueble_id": 10,
                "torre_apto": "431",
                "conjunto_id": 1,
                "conjunto_nombre": "COLON TEST",
                "conjunto_catalogo": "COLON TEST",
                "copropiedad_nombre": None,
                "titular_nombre": "A",
                "titular_cedula": "1",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
            {
                "inmueble_id": 11,
                "torre_apto": "1431",
                "conjunto_id": 1,
                "conjunto_nombre": "COLON TEST",
                "conjunto_catalogo": "COLON TEST",
                "copropiedad_nombre": None,
                "titular_nombre": "B",
                "titular_cedula": "2",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
            {
                "inmueble_id": 12,
                "torre_apto": "1-431",
                "conjunto_id": 1,
                "conjunto_nombre": "COLON TEST",
                "conjunto_catalogo": "COLON TEST",
                "copropiedad_nombre": None,
                "titular_nombre": "C",
                "titular_cedula": "3",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
        ]
        # Solo apto 431 (+ codigo 1431): rechaza 431, toma 1431 por código
        hallado = repo._elegir_por_unidad(
            filas,
            repo._claves_unidad("", bloque="", apartamento="431"),
            codigo_cuenta="1431",
        )
        self.assertIsNotNone(hallado)
        self.assertEqual(hallado["inmueble_id"], 11)
        self.assertEqual(hallado["torre_apto"], "1431")

        # Bloque+apto 1-431 + codigo 1431 + ambas filas: prioriza Identificación
        hallado2 = repo._elegir_por_unidad(
            filas,
            repo._claves_unidad("", bloque="1", apartamento="431"),
            codigo_cuenta="1431",
        )
        self.assertEqual(hallado2["inmueble_id"], 11)

        # Sin fila plana 1431: acepta estructura 1-431 (misma cuenta COLON)
        filas_sin_plano = [f for f in filas if f["inmueble_id"] != 11]
        hallado3 = repo._elegir_por_unidad(
            filas_sin_plano,
            repo._claves_unidad("", bloque="1", apartamento="431"),
            codigo_cuenta="1431",
        )
        self.assertEqual(hallado3["inmueble_id"], 12)

        # 1124 / 124 compacto
        filas_1124 = [
            {
                "inmueble_id": 20,
                "torre_apto": "124",
                "conjunto_id": 1,
                "conjunto_nombre": "X",
                "conjunto_catalogo": "X",
                "copropiedad_nombre": None,
                "titular_nombre": "Z",
                "titular_cedula": "9",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
            {
                "inmueble_id": 21,
                "torre_apto": "1124",
                "conjunto_id": 1,
                "conjunto_nombre": "X",
                "conjunto_catalogo": "X",
                "copropiedad_nombre": None,
                "titular_nombre": "Y",
                "titular_cedula": "8",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
        ]
        h1124 = repo._elegir_por_unidad(
            filas_1124,
            repo._claves_unidad("", bloque="1", apartamento="124"),
            codigo_cuenta="1124",
        )
        self.assertEqual(h1124["inmueble_id"], 21)

        # Torre+apto compacto texto: TORRE X APTO 1431 vs Neon 431
        filas_ta = [
            {
                "inmueble_id": 30,
                "torre_apto": "431",
                "conjunto_id": 1,
                "conjunto_nombre": "X",
                "conjunto_catalogo": "X",
                "copropiedad_nombre": None,
                "titular_nombre": "T",
                "titular_cedula": "7",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
            {
                "inmueble_id": 31,
                "torre_apto": "X-1431",
                "conjunto_id": 1,
                "conjunto_nombre": "X",
                "conjunto_catalogo": "X",
                "copropiedad_nombre": None,
                "titular_nombre": "U",
                "titular_cedula": "6",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
        ]
        h_ta = repo._elegir_por_unidad(
            filas_ta,
            repo._claves_unidad("TORRE X APTO 1431"),
            codigo_cuenta="1431",
        )
        self.assertEqual(h_ta["inmueble_id"], 31)
        # Sin match de código ni estructura compatible → None (no 431)
        h_bad = repo._elegir_por_unidad(
            [filas_ta[0]],
            repo._claves_unidad("", bloque="", apartamento="431"),
            codigo_cuenta="1431",
        )
        self.assertIsNone(h_bad)

    def test_score_conjunto_ignora_ph_y_acentos(self):
        self.assertGreaterEqual(
            repo._score_conjunto(
                "URBANIZACIÓN SANTA CLARA MANZANA 2 PH",
                "Urbanizacion Santa Clara Manzana 2",
            ),
            90,
        )
        self.assertGreaterEqual(
            repo._score_conjunto(
                "SANTA CLARA MANZANA 2",
                "URBANIZACIÓN SANTA CLARA MANZANA 2 PROPIEDAD HORIZONTAL",
            ),
            70,
        )
        self.assertEqual(
            repo._score_conjunto("OTRO CONJUNTO", "SANTA CLARA"),
            0,
        )

    def test_conjunto_prefijo_letra_y_etapa_mirador(self):
        """Regresión Mirador: 'D - … ETAPA 1 P.H.' ≡ 'MIRADOR DE LLANO GRANDE'."""
        pdf = "D - MIRADOR DE LLANO GRANDE ETAPA 1 P.H."
        neon = "MIRADOR DE LLANO GRANDE"
        self.assertEqual(
            repo._nucleo_conjunto(pdf),
            "MIRADOR DE LLANO GRANDE",
        )
        self.assertEqual(repo._nucleo_conjunto(neon), "MIRADOR DE LLANO GRANDE")
        self.assertGreaterEqual(repo._score_conjunto(neon, pdf), 90)
        # Variante con etapa en Neon también
        self.assertGreaterEqual(
            repo._score_conjunto("MIRADOR DE LLANO GRANDE ETAPA 1", pdf),
            90,
        )

    def test_buscar_mirador_02_042_en_memoria(self):
        filas = [
            {
                "inmueble_id": 42,
                "torre_apto": "2-42",
                "conjunto_id": 9,
                "conjunto_nombre": "MIRADOR DE LLANO GRANDE",
                "conjunto_catalogo": "MIRADOR DE LLANO GRANDE",
                "copropiedad_nombre": "MIRADOR DE LLANO GRANDE PH",
                "copropiedad_nit": "900",
                "ciudad": "Pereira",
                "titular_nombre": "MOSQUERA MOSQUERA LUZ ELVIRA",
                "titular_cedula": "1",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
            {
                "inmueble_id": 43,
                "torre_apto": "2-43",
                "conjunto_id": 9,
                "conjunto_nombre": "MIRADOR DE LLANO GRANDE",
                "conjunto_catalogo": "MIRADOR DE LLANO GRANDE",
                "copropiedad_nombre": "MIRADOR DE LLANO GRANDE PH",
                "copropiedad_nit": "900",
                "ciudad": "Pereira",
                "titular_nombre": "OTRO",
                "titular_cedula": "2",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
        ]
        filtradas = repo._filtrar_por_conjunto(
            filas,
            conjunto_id=None,
            conjunto_nombre="D - MIRADOR DE LLANO GRANDE ETAPA 1 P.H.",
        )
        self.assertEqual(len(filtradas), 2)
        claves = repo._claves_unidad("TORRE 02 APTO 042", "02", "042")
        hallado = repo._elegir_por_unidad(
            filtradas, claves, titular="MOSQUERA MOSQUERA LUZ ELVIRA"
        )
        self.assertIsNotNone(hallado)
        self.assertEqual(hallado["inmueble_id"], 42)
        self.assertEqual(hallado["torre_apto"], "2-42")

    def test_describir_busqueda_incluye_claves_y_codigo(self):
        txt = repo.describir_busqueda(
            conjunto_nombre="Demo PH",
            bloque="1",
            apartamento="201",
            torre_apto="TORRE 1 APTO 201",
            titular="ANA PEREZ",
            codigo_cuenta="9401",
        )
        self.assertIn("conjunto=", txt)
        self.assertIn("1-201", txt)
        self.assertIn("codigo_cuenta='9401'", txt)
        self.assertIn("Identificación COLON", txt)
        self.assertIn("anti-sufijo", txt)

    def test_buscar_match_torre_vs_guion_con_filas_en_memoria(self):
        filas = [
            {
                "inmueble_id": 10,
                "torre_apto": "1-201",
                "conjunto_id": 3,
                "conjunto_nombre": "SANTA CLARA MANZANA 2",
                "conjunto_catalogo": "SANTA CLARA MANZANA 2",
                "copropiedad_nombre": "SANTA CLARA MANZANA 2 PH",
                "copropiedad_nit": "900",
                "ciudad": "Pereira",
                "titular_nombre": "ANA PEREZ",
                "titular_cedula": "1",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
            {
                "inmueble_id": 11,
                "torre_apto": "1-202",
                "conjunto_id": 3,
                "conjunto_nombre": "SANTA CLARA MANZANA 2",
                "conjunto_catalogo": "SANTA CLARA MANZANA 2",
                "copropiedad_nombre": "SANTA CLARA MANZANA 2 PH",
                "copropiedad_nit": "900",
                "ciudad": "Pereira",
                "titular_nombre": "OTRO",
                "titular_cedula": "2",
                "contacto_nombre": None,
                "contacto_cedula": None,
            },
        ]
        filtradas = repo._filtrar_por_conjunto(
            filas,
            conjunto_id=None,
            conjunto_nombre="URBANIZACIÓN SANTA CLARA MANZANA 2",
        )
        self.assertEqual(len(filtradas), 2)
        claves = repo._claves_unidad("TORRE 1 APTO 201", "1", "201")
        hallado = repo._elegir_por_unidad(filtradas, claves)
        self.assertIsNotNone(hallado)
        self.assertEqual(hallado["inmueble_id"], 10)
        self.assertEqual(hallado["torre_apto"], "1-201")

    def test_elegir_desambigua_por_titular(self):
        filas = [
            {
                "inmueble_id": 1,
                "torre_apto": "9-401",
                "titular_nombre": "BANOL RIVERA MAURICIO",
                "titular_cedula": "1",
                "contacto_nombre": None,
                "contacto_cedula": None,
                "conjunto_nombre": "X",
                "conjunto_catalogo": "X",
                "copropiedad_nombre": "X",
                "copropiedad_nit": "1",
                "ciudad": None,
                "conjunto_id": 1,
            },
            {
                "inmueble_id": 2,
                "torre_apto": "9-401",
                "titular_nombre": "OTRO TITULAR",
                "titular_cedula": "2",
                "contacto_nombre": None,
                "contacto_cedula": None,
                "conjunto_nombre": "X",
                "conjunto_catalogo": "X",
                "copropiedad_nombre": "X",
                "copropiedad_nit": "1",
                "ciudad": None,
                "conjunto_id": 1,
            },
        ]
        claves = repo._claves_unidad("", "9", "401")
        hallado = repo._elegir_por_unidad(
            filas, claves, titular="BAÑOL RIVERA MAURICIO"
        )
        self.assertEqual(hallado["inmueble_id"], 1)


class ValidacionDatosTests(unittest.TestCase):
    def test_validar_reporta_faltantes_exactos(self):
        with self.assertRaises(svc.CertificadoDatosFaltantesError) as ctx:
            svc.validar_datos_criticos(
                {
                    "copropiedad_nombre": "",
                    "copropiedad_nit": "900",
                    "titular_nombre": "ANA",
                    "titular_cedula": "",
                    "torre_apto": "101",
                }
            )
        msg = str(ctx.exception)
        self.assertIn("nombre de la copropiedad", msg)
        self.assertIn("cédula del titular", msg)
        self.assertNotIn("NIT de la copropiedad", msg)
        self.assertEqual(len(ctx.exception.faltantes), 2)

    def test_validar_ok_cuando_completos(self):
        svc.validar_datos_criticos(
            {
                "copropiedad_nombre": "PH X",
                "copropiedad_nit": "900",
                "titular_nombre": "ANA",
                "titular_cedula": "1",
                "torre_apto": "T1-101",
            }
        )


class GeneracionWordTests(unittest.TestCase):
    def test_plantilla_oficial_versionada(self):
        self.assertTrue(svc.PLANTILLA_PATH.is_file())
        self.assertEqual(svc.PLANTILLA_PATH.name, "CERTIFICADO_DE_DEUDA.docx")

    def test_render_plantilla_oficial(self):
        ctx = svc.construir_contexto_plantilla(
            datos_neon={
                "copropiedad_nombre": "URBANIZACIÓN DEMO",
                "copropiedad_nit": "900123",
                "torre_apto": "TORRE 1 APTO 101",
                "conjunto_nombre": "URBANIZACIÓN DEMO",
                "titular_nombre": "ANA PEREZ",
                "titular_cedula": "123",
                "ciudad": "Pereira",
            },
            filas=svc.agrupar_capital_limpio(
                [
                    {
                        "fecha": "2024.02.01",
                        "concepto": "CUOTA ADMINISTRACION",
                        "valor_a_demandar": 100,
                    }
                ]
            ),
            representante_nombre="RL DEMO",
            representante_cedula="1.234.567",
        )
        self.assertIn("filas", ctx)
        self.assertEqual(ctx["filas"][0]["mes"], "FEBRERO")
        self.assertEqual(ctx["filas"][0]["anio"], "2024")
        self.assertEqual(ctx["total_saldo"], "$ 100")
        buf = svc.renderizar_docx(ctx)
        raw = buf.getvalue()
        self.assertTrue(raw[:2] == b"PK")
        with zipfile.ZipFile(BytesIO(raw)) as zf:
            self.assertIn("word/document.xml", zf.namelist())
            xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("ANA PEREZ", xml)
            self.assertIn("FEBRERO", xml)

    def test_plantillas_oficio_sin_parentesis_y_negritas(self):
        """Oficio 8.5×13, campos sin (( )), negritas y spacing single (after=0)."""
        from docx import Document
        from docx.oxml.ns import qn

        ctx = svc.construir_contexto_plantilla(
            datos_neon={
                "copropiedad_nombre": "MIRADOR DE LLANO GRANDE PH",
                "copropiedad_nit": "900938646",
                "torre_apto": "4-34",
                "conjunto_nombre": "MIRADOR DE LLANO GRANDE",
                "titular_nombre": "ALBA DOLLY CARDONA AMARILES",
                "titular_cedula": "21423830",
                "ciudad": "Pereira",
            },
            filas=svc.agrupar_capital_limpio(
                [
                    {
                        "fecha": "2026.02.01",
                        "concepto": "CUOTA ADMINISTRACION",
                        "valor_a_demandar": 30000,
                    }
                ]
            ),
            representante_nombre="GLADYS DEMO",
            representante_cedula="35.319.382",
        )
        doc = Document(svc.renderizar_docx(ctx))
        self.assertAlmostEqual(doc.sections[0].page_width.inches, 8.5, places=2)
        self.assertAlmostEqual(doc.sections[0].page_height.inches, 13.0, places=2)
        self.assertEqual(
            doc.paragraphs[1].text,
            "MIRADOR DE LLANO GRANDE PH PROPIEDAD HORIZONTAL",
        )
        self.assertEqual(doc.paragraphs[2].text, "N.I.T.: 900938646")
        self.assertNotIn("((", doc.paragraphs[1].text)
        self.assertNotIn("((", doc.paragraphs[2].text)
        cuerpo = doc.paragraphs[8].text
        self.assertIn("PROPIEDAD HORIZONTAL MIRADOR DE LLANO GRANDE,", cuerpo)
        self.assertIn("cedula 21423830,", cuerpo)
        self.assertNotIn("((", cuerpo)
        self.assertNotIn("))", cuerpo)
        self.assertIn("del (la) señor(a)", cuerpo)
        bold_cuerpo = "".join(r.text for r in doc.paragraphs[8].runs if r.bold)
        self.assertIn("ALBA DOLLY CARDONA AMARILES", bold_cuerpo)
        self.assertIn("21423830", bold_cuerpo)
        self.assertIn("PROPIEDAD HORIZONTAL MIRADOR DE LLANO GRANDE,", bold_cuerpo)
        sp = doc.paragraphs[8]._p.find(qn("w:pPr")).find(qn("w:spacing"))
        self.assertEqual(sp.get(qn("w:after")), "0")
        self.assertEqual(sp.get(qn("w:line")), "240")

        ctx_poder = svc.construir_contexto_poder(
            datos_neon={
                "copropiedad_nombre": "MIRADOR DE LLANO GRANDE PH",
                "copropiedad_nit": "900938646",
                "torre_apto": "4-34",
                "conjunto_nombre": "MIRADOR DE LLANO GRANDE",
                "titular_nombre": "ALBA DOLLY CARDONA AMARILES",
                "titular_cedula": "21423830",
                "ciudad": "Pereira",
            },
            capital_limpio_a_demandar=[
                {"fecha": "2026.02.01", "valor_a_demandar": 30000},
            ],
            representante_nombre="GLADYS DEMO",
            representante_cedula="35.319.382",
            fmi="290-204208",
        )
        docp = Document(svc.renderizar_poder(ctx_poder))
        self.assertAlmostEqual(docp.sections[0].page_height.inches, 13.0, places=2)
        cuerpo_p = next(p.text for p in docp.paragraphs if "CONFIERO PODER" in p.text)
        self.assertIn("propiedad horizontal MIRADOR DE LLANO GRANDE PH", cuerpo_p)
        self.assertIn("N.I.T.: 900938646", cuerpo_p)
        self.assertNotIn("((", cuerpo_p)
        self.assertNotIn("))", cuerpo_p)
        self.assertIn("identificado con el FMI 290-204208", cuerpo_p)
        bold_poder = "".join(
            r.text
            for p in docp.paragraphs
            if "CONFIERO PODER" in p.text
            for r in p.runs
            if r.bold
        )
        self.assertIn("MIRADOR DE LLANO GRANDE PH", bold_poder)
        self.assertIn("900938646", bold_poder)

    def test_plantillas_layout_enters_y_firmas_keep_with(self):
        """Single spacing + enters entre bloques; keepNext en antefirmas; 1 page break."""
        from docx import Document
        from docx.oxml.ns import qn

        def _has(p, tag: str) -> bool:
            pPr = p._p.find(qn("w:pPr"))
            return pPr is not None and pPr.find(qn(f"w:{tag}")) is not None

        def _chain_from(paragraphs, start_text: str, n: int):
            for i, p in enumerate(paragraphs):
                if p.text.strip() == start_text:
                    return paragraphs[i : i + n]
            self.fail(f"No se encontró párrafo {start_text!r}")

        neon = {
            "copropiedad_nombre": "MIRADOR DE LLANO GRANDE PH",
            "copropiedad_nit": "900938646",
            "torre_apto": "4-34",
            "conjunto_nombre": "MIRADOR DE LLANO GRANDE",
            "titular_nombre": "ALBA DOLLY CARDONA AMARILES",
            "titular_cedula": "21423830",
            "ciudad": "Pereira",
        }
        ctx = svc.construir_contexto_plantilla(
            datos_neon=neon,
            filas=svc.agrupar_capital_limpio(
                [
                    {
                        "fecha": "2026.02.01",
                        "concepto": "CUOTA ADMINISTRACION",
                        "valor_a_demandar": 30000,
                    }
                ]
            ),
            representante_nombre="GLADYS DEMO",
            representante_cedula="35.319.382",
        )
        doc = Document(svc.renderizar_docx(ctx))
        # Enters entre bloques del certificado (vacíos semánticos).
        self.assertEqual(doc.paragraphs[3].text.strip(), "")  # header↔juzgado
        self.assertEqual(doc.paragraphs[7].text.strip(), "")  # juzgado↔cuerpo
        self.assertEqual(doc.paragraphs[9].text.strip(), "")  # cuerpo↔tabla
        # Cadena antefirma CERT: Atentamente → 4 vacíos → nombre → C.C. keepNext; R.L sin.
        cert_chain = _chain_from(doc.paragraphs, "Atentamente,", 8)
        self.assertEqual(cert_chain[0].text.strip(), "Atentamente,")
        for p in cert_chain[1:5]:
            self.assertEqual(p.text.strip(), "")
        self.assertIn("GLADYS DEMO", cert_chain[5].text)
        self.assertIn("35.319.382", cert_chain[6].text)
        self.assertTrue(cert_chain[7].text.strip().startswith("R.L"))
        for p in cert_chain[:7]:
            self.assertTrue(_has(p, "keepNext"), p.text[:40])
            self.assertTrue(_has(p, "keepLines"), p.text[:40])
            self.assertTrue(_has(p, "widowControl"), p.text[:40])
        self.assertFalse(_has(cert_chain[7], "keepNext"))
        sp = doc.paragraphs[8]._p.find(qn("w:pPr")).find(qn("w:spacing"))
        self.assertEqual(sp.get(qn("w:line")), "240")
        self.assertEqual(sp.get(qn("w:after")), "0")

        ctx_poder = svc.construir_contexto_poder(
            datos_neon=neon,
            capital_limpio_a_demandar=[
                {"fecha": "2026.02.01", "valor_a_demandar": 30000},
            ],
            representante_nombre="GLADYS DEMO",
            representante_cedula="35.319.382",
            fmi="290-204208",
        )
        buf_poder = svc.renderizar_poder(ctx_poder)
        docp = Document(buf_poder)
        texts = [p.text.strip() for p in docp.paragraphs]
        mandato_i = next(i for i, t in enumerate(texts) if "CONFIERO PODER" in t)
        facultades_i = next(i for i, t in enumerate(texts) if "facultades expresas" in t)
        sirvase_i = next(i for i, t in enumerate(texts) if t.startswith("Sírvase"))
        self.assertEqual(texts[mandato_i + 1], "")
        self.assertEqual(facultades_i, mandato_i + 2)
        self.assertEqual(texts[facultades_i + 1], "")
        self.assertEqual(sirvase_i, facultades_i + 2)

        rl_chain = _chain_from(docp.paragraphs, "Atentamente,", 5)
        for p in rl_chain[:4]:
            self.assertTrue(_has(p, "keepNext"))
        self.assertFalse(_has(rl_chain[4], "keepNext"))
        self.assertIn("35.319.382", rl_chain[4].text)

        ab_chain = _chain_from(docp.paragraphs, "Acepto,", 7)
        for p in ab_chain[:6]:
            self.assertTrue(_has(p, "keepNext"))
        self.assertFalse(_has(ab_chain[6], "keepNext"))
        self.assertIn("L.T.", ab_chain[6].text)

        composed = svc.componer_certificado_con_poder(
            svc.renderizar_docx(ctx), buf_poder
        )
        with zipfile.ZipFile(composed) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        self.assertEqual(xml.count('w:type="page"'), 1)
        self.assertNotIn("pageBreakBefore", xml)

    def test_nombre_archivo_dinamico(self):
        self.assertEqual(
            svc.nombre_archivo_certificado("ANA PÉREZ GOMEZ"),
            "Certificado_ANA_PÉREZ_GOMEZ.docx",
        )

    def test_generar_falla_si_neon_no_encuentra(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value=None,
        ):
            with self.assertRaises(svc.CertificadoNoEncontradoError) as ctx:
                svc.generar_certificado_deuda(
                    capital_limpio_a_demandar=[],
                    conjunto="Demo",
                    bloque="1",
                    apartamento="201",
                    torre_apto="TORRE 1 APTO 201",
                    codigo_cuenta="9401",
                )
            msg = str(ctx.exception)
            self.assertIn("Buscado:", msg)
            self.assertIn("1-201", msg)
            self.assertIn("codigo_cuenta", msg)
            self.assertTrue(ctx.exception.criterios)

    def test_generar_falla_si_faltan_criticos(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value={
                "inmueble_id": 1,
                "copropiedad_nombre": "PH",
                "copropiedad_nit": "",
                "titular_nombre": "ANA",
                "titular_cedula": "1",
                "torre_apto": "101",
                "conjunto_nombre": "X",
            },
        ):
            with self.assertRaises(svc.CertificadoDatosFaltantesError) as ctx:
                svc.generar_certificado_deuda(
                    capital_limpio_a_demandar=[],
                    inmueble_id=1,
                )
            self.assertIn("NIT de la copropiedad", str(ctx.exception))

    def test_generar_ok_con_neon_mock(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value={
                "inmueble_id": 7,
                "copropiedad_nombre": "PH Demo",
                "copropiedad_nit": "900",
                "titular_nombre": "JUAN LOPEZ",
                "titular_cedula": "555",
                "torre_apto": "2-202",
                "conjunto_nombre": "Demo",
                "ciudad": "Pereira",
            },
        ):
            buf, nombre, meta = svc.generar_certificado_deuda(
                capital_limpio_a_demandar=[
                    {
                        "fecha": "2024.06.01",
                        "concepto": "CUOTA ADMINISTRACION",
                        "valor_a_demandar": 40,
                    },
                    {
                        "fecha": "2024.06.01",
                        "concepto": "EXTRAORDINARIA",
                        "valor_a_demandar": 10,
                    },
                ],
                inmueble_id=7,
                representante_nombre="RL",
                representante_cedula="99",
            )
        self.assertEqual(nombre, "Certificado_JUAN_LOPEZ.docx")
        self.assertEqual(meta["inmueble_id"], 7)
        self.assertEqual(meta["filas"], 1)
        self.assertEqual(meta["plantilla"], "CERTIFICADO_DE_DEUDA.docx")
        self.assertTrue(buf.getvalue()[:2] == b"PK")

    def test_generar_pasa_titular_al_resolver(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value={
                "inmueble_id": 3,
                "copropiedad_nombre": "PH",
                "copropiedad_nit": "1",
                "titular_nombre": "ANA",
                "titular_cedula": "9",
                "torre_apto": "1-101",
                "conjunto_nombre": "X",
                "ciudad": "Pereira",
            },
        ) as mocked:
            svc.generar_certificado_deuda(
                capital_limpio_a_demandar=[],
                conjunto="X",
                bloque="1",
                apartamento="101",
                titular="ANA PEREZ",
                representante_nombre="RL",
                representante_cedula="1",
            )
            kwargs = mocked.call_args.kwargs
            self.assertEqual(kwargs["titular"], "ANA PEREZ")
            self.assertEqual(kwargs["bloque"], "1")
            self.assertEqual(kwargs["apartamento"], "101")

    def test_permitir_datos_pdf_emite_sin_neon_si_hay_nit_cedula(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value=None,
        ):
            buf, nombre, meta = svc.generar_certificado_deuda(
                capital_limpio_a_demandar=[
                    {
                        "fecha": "2024.02.01",
                        "concepto": "CUOTA ADMINISTRACION",
                        "valor_a_demandar": 30,
                    }
                ],
                conjunto="D - MIRADOR DE LLANO GRANDE ETAPA 1 P.H.",
                bloque="02",
                apartamento="042",
                torre_apto="TORRE 02 APTO 042",
                titular="MOSQUERA MOSQUERA LUZ ELVIRA",
                titular_cedula="52.000.000",
                copropiedad_nit="900123456-1",
                permitir_datos_pdf=True,
                representante_nombre="RL",
                representante_cedula="1",
            )
        self.assertEqual(meta["fuente"], "pdf")
        self.assertIsNone(meta["inmueble_id"])
        self.assertIn("MOSQUERA", nombre)
        self.assertTrue(buf.getvalue()[:2] == b"PK")

    def test_permitir_datos_pdf_400_si_falta_nit_cedula(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value=None,
        ):
            with self.assertRaises(svc.CertificadoDatosFaltantesError) as ctx:
                svc.generar_certificado_deuda(
                    capital_limpio_a_demandar=[],
                    conjunto="MIRADOR",
                    bloque="02",
                    apartamento="042",
                    titular="MOSQUERA",
                    permitir_datos_pdf=True,
                )
            self.assertEqual(ctx.exception.fuente, "pdf")
            msg = str(ctx.exception)
            self.assertIn("sin inventar NIT/cédula", msg)
            self.assertIn("NIT", msg)
            self.assertIn("cédula", msg)

    def test_validar_antefirma_exige_nombre_y_cedula(self):
        with self.assertRaises(svc.CertificadoDatosFaltantesError) as ctx:
            svc.validar_antefirma_representante("", "1.2.3")
        self.assertEqual(ctx.exception.fuente, "antefirma")
        self.assertIn("nombre completo", str(ctx.exception))

        with self.assertRaises(svc.CertificadoDatosFaltantesError) as ctx2:
            svc.validar_antefirma_representante("RL DEMO", "  ")
        self.assertEqual(ctx2.exception.fuente, "antefirma")
        self.assertIn("cédula", str(ctx2.exception))

        nombre, cedula = svc.validar_antefirma_representante(
            "  GLADYS DEMO  ", " 35.319.382 "
        )
        self.assertEqual(nombre, "GLADYS DEMO")
        self.assertEqual(cedula, "35.319.382")

    def test_generar_falla_sin_antefirma_representante(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value={
                "inmueble_id": 7,
                "copropiedad_nombre": "PH Demo",
                "copropiedad_nit": "900",
                "titular_nombre": "JUAN LOPEZ",
                "titular_cedula": "555",
                "torre_apto": "2-202",
                "conjunto_nombre": "Demo",
                "ciudad": "Pereira",
            },
        ):
            with self.assertRaises(svc.CertificadoDatosFaltantesError) as ctx:
                svc.generar_certificado_deuda(
                    capital_limpio_a_demandar=[
                        {
                            "fecha": "2024.06.01",
                            "concepto": "CUOTA ADMINISTRACION",
                            "valor_a_demandar": 40,
                        }
                    ],
                    inmueble_id=7,
                )
            self.assertEqual(ctx.exception.fuente, "antefirma")
            self.assertIn("representante", str(ctx.exception).lower())

    def test_contexto_incluye_representante_en_placeholders(self):
        ctx = svc.construir_contexto_plantilla(
            datos_neon={
                "copropiedad_nombre": "PH",
                "copropiedad_nit": "1",
                "torre_apto": "1-1",
                "conjunto_nombre": "PH",
                "titular_nombre": "A",
                "titular_cedula": "2",
                "ciudad": "Pereira",
            },
            filas=[],
            representante_nombre="RL FIRMA",
            representante_cedula="99.888.777",
        )
        self.assertEqual(ctx["representante_nombre"], "RL FIRMA")
        self.assertEqual(ctx["representante_cedula"], "99.888.777")
        self.assertIn("representante_nombre", svc.PLANTILLA_PLACEHOLDERS)
        self.assertIn("representante_cedula", svc.PLANTILLA_PLACEHOLDERS)


class PoderConCertificadoTests(unittest.TestCase):
    def test_plantilla_poder_versionada(self):
        self.assertTrue(svc.PLANTILLA_PODER_PATH.is_file())
        self.assertEqual(svc.PLANTILLA_PODER_PATH.name, "PODER.docx")

    def test_periodo_desde_capital(self):
        desde, hasta = svc.periodo_desde_capital(
            [
                {"fecha": "2024.02.01", "valor_a_demandar": 1},
                {"fecha": "2024.06.15", "valor_a_demandar": 1},
            ]
        )
        self.assertEqual(desde, "febrero de 2024")
        self.assertEqual(hasta, "junio de 2024")
        vacio_d, vacio_h = svc.periodo_desde_capital([])
        self.assertEqual(vacio_d, "________")
        self.assertEqual(vacio_h, "________")

    def test_contexto_poder_no_inventa_fmi(self):
        ctx = svc.construir_contexto_poder(
            datos_neon={
                "copropiedad_nombre": "PH Demo",
                "copropiedad_nit": "900",
                "torre_apto": "2-202",
                "conjunto_nombre": "Demo",
                "titular_nombre": "JUAN LOPEZ",
                "titular_cedula": "555",
                "ciudad": "Pereira",
                "copropiedad_direccion": "Carrera 1",
                "copropiedad_telefono": "300",
                "copropiedad_email": "a@b.com",
            },
            capital_limpio_a_demandar=[
                {"fecha": "2024.06.01", "valor_a_demandar": 40},
            ],
            representante_nombre="GLADYS DEMO",
            representante_cedula="35.319.382",
        )
        self.assertEqual(ctx["fmi"], "")
        self.assertNotIn("copropiedad_direccion", ctx)
        self.assertNotIn("resolucion_numero", ctx)
        self.assertNotIn("representante_cedula_expedida_en", ctx)
        self.assertEqual(ctx["ejecutante_nit"], "900")
        self.assertIn("JUAN LOPEZ", ctx["ejecutados_texto"])
        self.assertEqual(ctx["apoderado_nombre"], svc.APODERADO_NOMBRE)
        self.assertEqual(ctx["periodo_desde"], "junio de 2024")
        buf = svc.renderizar_poder(ctx)
        self.assertTrue(buf.getvalue()[:2] == b"PK")
        with zipfile.ZipFile(BytesIO(buf.getvalue())) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("JUAN LOPEZ", xml)
            self.assertIn(svc.APODERADO_NOMBRE, xml)
            self.assertNotIn("290-219335", xml)
            self.assertNotIn("Carrera 1", xml)
            self.assertNotIn("Resolución número ________", xml)
            self.assertNotIn("expedida en la ciudad", xml)

    def test_contexto_poder_incluye_fmi_si_se_pasa(self):
        ctx = svc.construir_contexto_poder(
            datos_neon={
                "copropiedad_nombre": "PH Demo",
                "copropiedad_nit": "900",
                "torre_apto": "2-202",
                "conjunto_nombre": "Demo",
                "titular_nombre": "JUAN LOPEZ",
                "titular_cedula": "555",
                "ciudad": "Pereira",
            },
            capital_limpio_a_demandar=[
                {"fecha": "2024.06.01", "valor_a_demandar": 40},
            ],
            representante_nombre="GLADYS DEMO",
            representante_cedula="35.319.382",
            fmi="290-219335",
        )
        self.assertEqual(ctx["fmi"], "290-219335")
        buf = svc.renderizar_poder(ctx)
        with zipfile.ZipFile(BytesIO(buf.getvalue())) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("290-219335", xml)
            self.assertIn("identificado con el FMI", xml)

    def test_generar_incluye_poder_por_default(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value={
                "inmueble_id": 7,
                "copropiedad_nombre": "PH Demo",
                "copropiedad_nit": "900",
                "titular_nombre": "JUAN LOPEZ",
                "titular_cedula": "555",
                "torre_apto": "2-202",
                "conjunto_nombre": "Demo",
                "ciudad": "Pereira",
            },
        ):
            buf, nombre, meta = svc.generar_certificado_deuda(
                capital_limpio_a_demandar=[
                    {
                        "fecha": "2024.06.01",
                        "concepto": "CUOTA ADMINISTRACION",
                        "valor_a_demandar": 40,
                    }
                ],
                inmueble_id=7,
                representante_nombre="RL",
                representante_cedula="99",
            )
        self.assertEqual(nombre, "Certificado_JUAN_LOPEZ.docx")
        self.assertTrue(meta["incluir_poder"])
        self.assertEqual(meta["plantilla_poder"], "PODER.docx")
        with zipfile.ZipFile(BytesIO(buf.getvalue())) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("CERTIFICADO DE DEUDA", xml)
            self.assertIn("Asunto", xml)
            self.assertIn("Poder", xml)
            self.assertIn("Acepto", xml)
            self.assertIn(svc.APODERADO_NOMBRE, xml)

    def test_generar_sin_poder_si_flag_false(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value={
                "inmueble_id": 7,
                "copropiedad_nombre": "PH Demo",
                "copropiedad_nit": "900",
                "titular_nombre": "JUAN LOPEZ",
                "titular_cedula": "555",
                "torre_apto": "2-202",
                "conjunto_nombre": "Demo",
                "ciudad": "Pereira",
            },
        ):
            buf, _nombre, meta = svc.generar_certificado_deuda(
                capital_limpio_a_demandar=[
                    {
                        "fecha": "2024.06.01",
                        "concepto": "CUOTA ADMINISTRACION",
                        "valor_a_demandar": 40,
                    }
                ],
                inmueble_id=7,
                representante_nombre="RL",
                representante_cedula="99",
                incluir_poder=False,
            )
        self.assertFalse(meta.get("incluir_poder"))
        self.assertNotIn("plantilla_poder", meta)
        with zipfile.ZipFile(BytesIO(buf.getvalue())) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("CERTIFICADO DE DEUDA", xml)
            self.assertNotIn(svc.APODERADO_NOMBRE, xml)
            self.assertNotIn("Acepto", xml)

    def test_generar_poder_con_fmi_en_meta(self):
        with patch(
            "certificados_deuda_service.resolver_datos_certificado",
            return_value={
                "inmueble_id": 7,
                "copropiedad_nombre": "PH Demo",
                "copropiedad_nit": "900",
                "titular_nombre": "JUAN LOPEZ",
                "titular_cedula": "555",
                "torre_apto": "2-202",
                "conjunto_nombre": "Demo",
                "ciudad": "Pereira",
            },
        ):
            buf, _nombre, meta = svc.generar_certificado_deuda(
                capital_limpio_a_demandar=[
                    {
                        "fecha": "2024.06.01",
                        "concepto": "CUOTA ADMINISTRACION",
                        "valor_a_demandar": 40,
                    }
                ],
                inmueble_id=7,
                representante_nombre="RL",
                representante_cedula="99",
                fmi="290-999",
            )
        self.assertEqual(meta.get("fmi"), "290-999")
        with zipfile.ZipFile(BytesIO(buf.getvalue())) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("290-999", xml)


class RbacCertificadoTests(unittest.TestCase):
    def test_post_exige_accion_editar(self):
        consulta = permisos.permisos_de_perfil(permisos.PERFIL_CONSULTA)
        abogado = permisos.permisos_de_perfil(permisos.PERFIL_ABOGADO)
        path = "/herramientas/estado-cuenta/certificado-deuda"
        self.assertTrue(permisos.denegar_acceso(consulta, "POST", path))
        self.assertFalse(permisos.denegar_acceso(abogado, "POST", path))


class FmiPorIndiceParseTests(unittest.TestCase):
    def test_parse_objeto(self):
        from certificados_deuda_router import _parse_fmi_por_indice

        self.assertEqual(
            _parse_fmi_por_indice('{"0":"290-1","2":" 290-2 "}'),
            {0: "290-1", 2: "290-2"},
        )
        self.assertIsNone(_parse_fmi_por_indice(""))
        self.assertEqual(
            _parse_fmi_por_indice('[[1,"A"],{"indice":3,"fmi":"B"}]'),
            {1: "A", 3: "B"},
        )


if __name__ == "__main__":
    unittest.main()
