import unittest

import pandas as pd

from estado_cuenta_etl import (
    CLASIFICACION_CUOTA,
    CLASIFICACION_EXTRA,
    anotar_verificacion,
    clasificar_concepto,
    depurar_movimientos,
)


def _fila(fecha, concepto, tipo, valor, abono, saldo=0, titular="ANA", archivo="a.pdf", numero="1"):
    return {
        "Archivo": archivo,
        "Titular": titular,
        "Bloque": "9",
        "Apartamento": "401",
        "Codigo Cuenta": "9401",
        "Concepto": concepto,
        "Tipo Documento": tipo,
        "Número": numero,
        "Fecha": fecha,
        "Valor": valor,
        "Abono": abono,
        "Saldo": saldo,
    }


class VerificacionTests(unittest.TestCase):
    def test_columnas_siguen_el_abono_desde_la_ultima_fila(self):
        df = pd.DataFrame(
            [
                _fila("2016.05.01", "CUOTA DE ADMINISTRACION", "FAC", 25000, 0, saldo=25000, numero="1"),
                _fila("2016.05.17", "CUOTA DE ADMINISTRACION", "RDC", 0, 25000, saldo=0, numero="2"),
                _fila("2016.06.01", "CUOTA DE ADMINISTRACION", "FAC", 25000, 0, saldo=25000, numero="3"),
            ]
        )
        salida = anotar_verificacion(df)
        self.assertEqual(salida["Abono Acumulado"].tolist(), [25000.0, 25000.0, 0.0])
        self.assertEqual(salida["Diferencia"].tolist(), [0.0, -25000.0, 25000.0])

    def test_hoja_deudor_va_de_concepto_a_la_formula(self):
        from openpyxl import load_workbook

        from estado_cuenta_pdf_service import generar_excel_lote

        filas = [
            _fila("2016.05.01", "CUOTA DE ADMINISTRACION", "FAC", 25000, 0, saldo=25000, numero="0000008"),
            _fila("2016.05.17", "INTERESES DE MORA", "FAC", 3000, 0, saldo=28000, numero="0000011"),
            _fila("2016.05.17", "CUOTA DE ADMINISTRACION", "RDC", 0, 25000, saldo=0, numero="0000023"),
        ]
        cuentas = [
            {
                "archivo": "a.pdf",
                "titular": "ANA",
                "bloque": "9",
                "apartamento": "401",
                "codigo_cuenta": "9401",
                "conjunto": "",
                "fechas_detectadas": 2,
                "movimientos_extraidos": 2,
                "bloques_omitidos": 0,
                "inconsistencias_saldo": 0,
                "saldo_final": 0,
                "rows": filas,
                "error": "",
            }
        ]
        libro = load_workbook(generar_excel_lote(cuentas))
        hoja = libro["401 ANA"]
        self.assertEqual(hoja["A1"].value, "Titular")
        self.assertEqual(hoja["B1"].value, "ANA")
        self.assertEqual(hoja["A5"].value, "Inicio mora")
        self.assertEqual(
            [celda.value for celda in hoja[7]],
            [
                "Concepto",
                "Tipo Documento",
                "Número",
                "Fecha",
                "Valor",
                "Abono",
                "Saldo",
                "Abono Acumulado",
                "Diferencia",
            ],
        )
        conceptos = [hoja.cell(fila, 1).value for fila in range(8, hoja.max_row + 1)]
        tipos = [hoja.cell(fila, 2).value for fila in range(8, hoja.max_row + 1)]
        self.assertEqual(conceptos, ["CUOTA DE ADMINISTRACION"])
        self.assertEqual(tipos, ["FAC"])
        movimientos = libro["Movimientos"]
        tipos_mov = [movimientos.cell(fila, 7).value for fila in range(2, movimientos.max_row + 1)]
        self.assertEqual(tipos_mov, ["FAC"])

    def test_la_verificacion_no_mezcla_cuentas(self):
        filas = [
            _fila("2016.05.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 0, saldo=10),
            _fila("2016.06.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 5, saldo=15, titular="OTRO", archivo="b.pdf"),
        ]
        salida = anotar_verificacion(pd.DataFrame(filas))
        self.assertEqual(salida["Abono Acumulado"].tolist(), [0.0, 5.0])
        self.assertEqual(salida["Diferencia"].tolist(), [10.0, 10.0])


class ClasificacionTests(unittest.TestCase):
    def test_typos_de_cuota_no_alteran_el_texto(self):
        for concepto in (
            "CUOTA DE ADMINISTRACION",
            "CUOTA ADMINISTRATIVA",
            "COUTA DE ADMINISTRACION",
            "APORTE CUOTA DE ADMINISTRACION",
        ):
            self.assertEqual(clasificar_concepto(concepto), CLASIFICACION_CUOTA)
        self.assertEqual(clasificar_concepto("SANCION ASAMBLEA"), CLASIFICACION_EXTRA)
        self.assertEqual(clasificar_concepto("INTERESES DE MORA"), "INTERES")


class CorteMoraTests(unittest.TestCase):
    def _base(self):
        # M acumula abonos desde la última fila. N = Saldo PDF - M.
        # Negativos desde abajo: 2024-07 (1.º), 2024-05 (2.º), 2024-03 (3.º) y 2024-01 (4.º).
        return [
            _fila("2024.01.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 0, saldo=0, numero="0"),
            _fila("2024.03.01", "SANCION ASAMBLEA", "FAC", 10, 80, saldo=305, numero="1"),
            _fila("2024.03.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 0, saldo=0, numero="2"),
            _fila("2024.04.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 60, saldo=225, numero="3"),
            _fila("2024.05.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 0, saldo=0, numero="4"),
            _fila("2024.06.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 40, saldo=165, numero="5"),
            _fila("2024.07.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 0, saldo=0, numero="6"),
            _fila("2024.08.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 20, saldo=125, numero="7"),
            _fila("2024.08.01", "INTERESES DE MORA", "FAC", 0, 0, saldo=105, numero="8"),
            _fila("2024.08.01", "CUOTA DE ADMINISTRACION", "RDC", 0, 5, saldo=105, numero="9"),
        ]

    def test_conserva_el_tercer_negativo_y_el_dia_completo(self):
        resultado = depurar_movimientos(pd.DataFrame(self._base()))
        fechas = resultado["depurado"]["Fecha"].tolist()
        conceptos = resultado["depurado"]["Concepto"].tolist()

        self.assertNotIn("2024-01-01", fechas)
        self.assertIn("2024-03-01", fechas)
        self.assertEqual(resultado["cortes"].iloc[0]["Fecha Inicio Mora"], "2024-03-01")
        self.assertIn("SANCION ASAMBLEA", conceptos)
        self.assertNotIn("INTERESES DE MORA", conceptos)
        self.assertTrue((resultado["depurado"]["Tipo Documento"] == "FAC").all())
        self.assertFalse(resultado["depurado"]["Concepto"].astype(str).str.contains("SÓLO EXTRAS").any())
        sancion = resultado["depurado"].loc[resultado["depurado"]["Concepto"] == "SANCION ASAMBLEA"].iloc[0]
        self.assertEqual(sancion["Clasificacion"], CLASIFICACION_EXTRA)
        self.assertEqual(sancion["Concepto"], "SANCION ASAMBLEA")

    def test_huerfano_conserva_el_concepto_original(self):
        df = pd.DataFrame(
            [
                _fila("2024.01.01", "SANCION ASAMBLEA", "FAC", 10, 0),
                _fila("2024.02.01", "ELABORACION PISO", "FAC", 10, 0),
            ]
        )
        resultado = depurar_movimientos(df)
        self.assertEqual(resultado["depurado"]["Concepto"].tolist(), ["SANCION ASAMBLEA", "ELABORACION PISO"])
        self.assertFalse(resultado["cortes"].iloc[0]["Corte Aplicado"])

    def test_no_mezcla_titulares(self):
        filas = self._base() + [_fila("2020.01.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 0, titular="OTRO", archivo="b.pdf")]
        resultado = depurar_movimientos(pd.DataFrame(filas))
        otro = resultado["depurado"].loc[resultado["depurado"]["Titular"] == "OTRO", "Fecha"].tolist()
        self.assertEqual(otro, ["2020-01-01"])
        self.assertEqual(len(resultado["certificados"]), 2)

    def test_celdas_rotas_no_cuentan_como_negativo(self):
        df = pd.DataFrame(
            [
                _fila("2024.01.01", "CUOTA DE ADMINISTRACION", "FAC", "#¡VALOR!", "########"),
                _fila("2024.02.01", "CUOTA DE ADMINISTRACION", "FAC", 10, 0),
            ]
        )
        resultado = depurar_movimientos(df)
        self.assertFalse(resultado["cortes"].iloc[0]["Corte Aplicado"])
        self.assertIn("valor_ilegible", resultado["depurado"].iloc[0]["Alerta"])

    def test_un_certificado_por_deudor_con_saldo_corrido(self):
        resultado = depurar_movimientos(pd.DataFrame(self._base()))
        hojas = {c["hoja"]: c for c in resultado["certificados"]}
        self.assertEqual(len(hojas), 1)
        tabla = next(iter(hojas.values()))["tabla"]
        marzo = tabla.loc[tabla["MES"] == "MARZO"].iloc[0]
        agosto = tabla.loc[tabla["MES"] == "AGOSTO"].iloc[0]
        self.assertEqual(marzo["CUOTAS ORDINARIAS"], 10)
        self.assertEqual(marzo["CUOTAS EXTRAORDINARIAS"], 10)
        # Último mes = solo su cobro (10). Marzo suma ese saldo de abajo hacia arriba: 70.
        self.assertEqual(agosto["SALDO"], 10)
        self.assertEqual(marzo["SALDO"], 70)
        self.assertTrue(pd.isna(tabla.loc[tabla["MES"] == "ABRIL", "CUOTAS EXTRAORDINARIAS"].iloc[0]))
        self.assertNotIn("2024-01-01", tabla["FECHA CAUSACION"].astype(str).tolist())

    def test_consolidado_suma_sin_borrar_conceptos(self):
        resultado = depurar_movimientos(pd.DataFrame(self._base()))
        marzo = resultado["consolidado"].loc[resultado["consolidado"]["Periodo"] == "2024-03"].iloc[0]
        self.assertEqual(marzo["Valor Cuota"], 10)
        self.assertEqual(marzo["Valor Extras"], 10)
        self.assertIn("SANCION ASAMBLEA", marzo["Conceptos Originales"])
        self.assertIn("CUOTA DE ADMINISTRACION", marzo["Conceptos Originales"])


class CacheLoteTests(unittest.TestCase):
    def test_analizar_no_arma_excel_y_la_descarga_lo_reusa(self):
        from estado_cuenta_pdf_service import excel_desde_cache, procesar_lote_estados_cuenta

        resultado = procesar_lote_estados_cuenta(
            [("malo.pdf", b"esto no es un pdf")],
            incluir_excel=False,
        )
        self.assertIsNone(resultado["excel"])
        self.assertTrue(resultado["cache_id"])
        self.assertEqual(resultado["archivos_fallidos"], 1)
        libro = excel_desde_cache(resultado["cache_id"])
        self.assertIsNotNone(libro)
        self.assertTrue(libro.getvalue().startswith(b"PK"))
        otra = excel_desde_cache(resultado["cache_id"])
        self.assertTrue(otra.getvalue().startswith(b"PK"))
        self.assertIsNone(excel_desde_cache("no-existe"))


if __name__ == "__main__":
    unittest.main()
