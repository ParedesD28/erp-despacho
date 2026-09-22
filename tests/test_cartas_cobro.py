"""Tests unitarios del módulo de cartas de cobro Word."""
from __future__ import annotations

import sys
import unittest
import zipfile
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import cartas_cobro_service as svc


class CartasCobroUnitTests(unittest.TestCase):
    def test_fecha_limite_default_mas_cinco_dias(self):
        self.assertEqual(svc.fecha_limite_default(date(2026, 9, 22)), date(2026, 9, 27))

    def test_dedupe_por_inmueble_conserva_primero(self):
        filas = [
            {"inmueble_id": 1, "deudor_nombre": "A"},
            {"inmueble_id": 1, "deudor_nombre": "B"},
            {"inmueble_id": 2, "deudor_nombre": "C"},
        ]
        out = svc.dedupe_por_inmueble(filas)
        self.assertEqual(len(out), 2)
        self.assertEqual(out[0]["deudor_nombre"], "A")
        self.assertEqual(out[1]["inmueble_id"], 2)

    def test_resolver_seleccion_rechaza_duplicados_y_faltantes(self):
        candidatos = [
            {"inmueble_id": 10, "saldo_total": 1000},
            {"inmueble_id": 20, "saldo_total": 2000},
        ]
        sel = svc.resolver_seleccion(candidatos, [20, 20, 10])
        self.assertEqual([c["inmueble_id"] for c in sel], [20, 10])
        with self.assertRaises(ValueError):
            svc.resolver_seleccion(candidatos, [99])
        with self.assertRaises(ValueError):
            svc.resolver_seleccion(candidatos, [])

    def test_docx_incluye_campos_y_atn_codeudor(self):
        candidato = {
            "inmueble_id": 7,
            "deudor_nombre": "ANA PEREZ",
            "deudor_identificacion": "123456",
            "torre_apto": "T1-101",
            "conjunto_residencial": "Altos del Parque",
            "nombre_ph": "PH Altos del Parque",
            "saldo_total": 1500000,
            "codeudor_nombre": "JUAN PEREZ",
        }
        data = svc.construir_carta_docx(
            candidato,
            fecha_limite=date(2026, 9, 27),
            fecha_carta=date(2026, 9, 22),
        )
        self.assertTrue(data.startswith(b"PK"))
        # Lectura vía python-docx
        from docx import Document

        doc = Document(BytesIO(data))
        texto = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("ANA PEREZ", texto)
        self.assertIn("C.C. 123456", texto)
        self.assertIn("T1-101", texto)
        self.assertIn("Altos del Parque", texto)
        self.assertIn("Atn. JUAN PEREZ", texto)
        self.assertIn("REQUERIMIENTO DE PAGO PREJURÍDICO - PH Altos del Parque", texto)
        self.assertIn("310 6927812", texto)
        self.assertIn("notificacionesdiegoparedes@outlook.com", texto)
        self.assertIn("Diego Alejandro Paredes García", texto)
        self.assertIn("$ 1.500.000", texto)
        self.assertIn("27 de septiembre de 2026", texto)

    def test_docx_sin_codeudor_no_incluye_atn(self):
        candidato = {
            "inmueble_id": 8,
            "deudor_nombre": "SOLO DEUDOR",
            "deudor_identificacion": "999",
            "torre_apto": "A-2",
            "conjunto_residencial": "Conjunto X",
            "nombre_ph": "PH X",
            "saldo_total": 100,
            "codeudor_nombre": None,
        }
        from docx import Document

        doc = Document(
            BytesIO(
                svc.construir_carta_docx(
                    candidato,
                    fecha_limite=date(2026, 10, 1),
                    fecha_carta=date(2026, 9, 22),
                )
            )
        )
        texto = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("Atn.", texto)

    def test_paquete_zip_varias_cartas(self):
        items = [
            {
                "inmueble_id": 1,
                "deudor_nombre": "A",
                "deudor_identificacion": "1",
                "torre_apto": "101",
                "conjunto_residencial": "Demo",
                "nombre_ph": "PH Demo",
                "saldo_total": 10,
            },
            {
                "inmueble_id": 2,
                "deudor_nombre": "B",
                "deudor_identificacion": "2",
                "torre_apto": "102",
                "conjunto_residencial": "Demo",
                "nombre_ph": "PH Demo",
                "saldo_total": 20,
            },
        ]
        contenido, filename, media = svc.generar_paquete_docx(
            items, fecha_limite=date(2026, 9, 30), fecha_carta=date(2026, 9, 22)
        )
        self.assertEqual(media, "application/zip")
        self.assertTrue(filename.endswith(".zip"))
        with zipfile.ZipFile(BytesIO(contenido)) as zf:
            names = zf.namelist()
            self.assertEqual(len(names), 2)
            self.assertTrue(all(n.endswith(".docx") for n in names))

    def test_listar_candidatos_excluye_saldo_no_verificado(self):
        filas = [
            {
                "inmueble_id": 1,
                "obligacion_id": 11,
                "deudor_nombre": "OK",
                "deudor_identificacion": "1",
                "torre_apto": "1",
                "conjunto_residencial": "C",
                "nombre_ph": "PH",
                "nit_ph": "",
                "tipo_cartera": "PREJURIDICO",
                "radicado_interno": None,
                "codeudor_nombre": None,
                "codeudor_identificacion": None,
                "conjunto_id": 5,
                "deudor_contacto_id": 9,
            },
            {
                "inmueble_id": 2,
                "obligacion_id": 22,
                "deudor_nombre": "BAD",
                "deudor_identificacion": "2",
                "torre_apto": "2",
                "conjunto_residencial": "C",
                "nombre_ph": "PH",
                "nit_ph": "",
                "tipo_cartera": "PREJURIDICO",
                "radicado_interno": None,
                "codeudor_nombre": None,
                "codeudor_identificacion": None,
                "conjunto_id": 5,
                "deudor_contacto_id": 8,
            },
        ]

        class FakeCur:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def execute(self, *args, **kwargs):
                return None

            def fetchall(self):
                return filas

        class FakeConn:
            def cursor(self, cursor_factory=None):
                return FakeCur()

            def release(self):
                return None

        def fake_saldo(oid, fecha_corte=None):
            if oid == 11:
                return {"saldo_verificado": True, "saldo_total": 5000, "saldo_fuente": "LIQUIDADOR_PH"}
            return {"saldo_verificado": False, "saldo_total": None, "saldo_fuente": "ERROR"}

        with patch.object(svc.db, "get_connection", return_value=FakeConn()), patch.object(
            svc.catalogos_service,
            "obtener_conjunto",
            return_value={"id": 5, "nombre": "C", "persona_juridica": "PH", "nit": "900"},
        ), patch.object(
            svc.obligacion_saldo_service, "calcular_saldo_obligacion", side_effect=fake_saldo
        ):
            out = svc.listar_candidatos(conjunto_id=5, fecha_corte=date(2026, 9, 22))
        self.assertEqual(len(out), 1)
        self.assertEqual(out[0]["inmueble_id"], 1)
        self.assertEqual(out[0]["saldo_total"], 5000.0)


if __name__ == "__main__":
    unittest.main()
