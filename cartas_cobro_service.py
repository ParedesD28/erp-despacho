"""Cartas de cobro prejurídico en Word (.docx).

Lista candidatos PH por conjunto, deduplica por inmueble, genera documentos
con python-docx. El saldo proviene únicamente de obligacion_saldo_service
(gran_total del liquidador PH). No modifica liquidador ni SMS.
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime, timedelta
from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt, Cm
from psycopg2.extras import RealDictCursor

import catalogos_service
import db
import obligacion_saldo_service

TELEFONO_DESPACHO = "310 6927812"
CORREO_DESPACHO = "notificacionesdiegoparedes@outlook.com"
FIRMANTE_NOMBRE = "Diego Alejandro Paredes García"
FIRMANTE_CARGO = "Abogado Apoderado"
CIUDAD = "Pereira"
DIAS_LIMITE_DEFAULT = 5
MAX_LOTE = 100
CARTERAS_VALIDAS = {"PREJURIDICO", "JURIDICO"}


def fecha_limite_default(hoy: Optional[date] = None) -> date:
    return (hoy or date.today()) + timedelta(days=DIAS_LIMITE_DEFAULT)


def _fmt_money(value: Any) -> str:
    try:
        return f"$ {float(value or 0):,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return "$ 0"


def _fmt_fecha_es(value: date) -> str:
    meses = (
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
    )
    return f"{value.day} de {meses[value.month - 1]} de {value.year}"


def _safe_filename(text: str) -> str:
    cleaned = re.sub(r"[^\w\-]+", "_", (text or "").strip(), flags=re.UNICODE)
    return cleaned.strip("_")[:80] or "sin_nombre"


def listar_candidatos(
    *,
    conjunto_id: int,
    tipo_cartera: str = "PREJURIDICO",
    fecha_corte: Optional[date] = None,
    saldo_minimo: float = 0.0,
    conn=None,
) -> list[dict]:
    """Candidatos PH del conjunto con saldo verificado, uno por inmueble."""
    cartera = (tipo_cartera or "PREJURIDICO").upper().strip()
    if cartera and cartera not in CARTERAS_VALIDAS:
        raise ValueError("Tipo de cartera no válido.")

    owns_conn = conn is None
    if owns_conn:
        conn = db.get_connection()

    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            conjunto = catalogos_service.obtener_conjunto(cur, int(conjunto_id))
            if not conjunto:
                raise ValueError("Conjunto no encontrado.")

            where = [
                "o.inmueble_id IS NOT NULL",
                "UPPER(COALESCE(o.fuente_saldo,''))='EXPENSAS_PH'",
                "UPPER(COALESCE(o.estado,'ACTIVA')) NOT IN ('CANCELADA','ANULADA','PAGADA','INACTIVO','INACTIVA')",
                "NOT EXISTS ("
                " SELECT 1 FROM proceso_obligaciones po_inact"
                " JOIN procesos p_inact ON p_inact.radicado_interno=po_inact.radicado_interno"
                " WHERE po_inact.obligacion_id=o.id"
                "   AND UPPER(COALESCE(p_inact.estado,'ACTIVO'))='INACTIVO'"
                ")",
                "(i.conjunto_id=%s OR ("
                " i.conjunto_id IS NULL"
                " AND UPPER(BTRIM(COALESCE(i.conjunto_residencial,'')))=UPPER(BTRIM(%s))"
                "))",
            ]
            params: list[Any] = [int(conjunto_id), conjunto["nombre"]]

            if cartera:
                # Sin proceso vinculado se trata como prejurídico (mismo criterio SMS).
                where.append(
                    "(COALESCE(p.tipo_cartera,'PREJURIDICO')=%s OR p.radicado_interno IS NULL)"
                    if cartera == "PREJURIDICO"
                    else "COALESCE(p.tipo_cartera,'PREJURIDICO')=%s"
                )
                params.append(cartera)

            cur.execute(
                f"""
                SELECT DISTINCT ON (i.id)
                    i.id AS inmueble_id,
                    i.conjunto_id,
                    COALESCE(i.conjunto_residencial, %s) AS conjunto_residencial,
                    i.torre_apto,
                    o.id AS obligacion_id,
                    deudor.contacto_id AS deudor_contacto_id,
                    deudor.nombre AS deudor_nombre,
                    deudor.identificacion AS deudor_identificacion,
                    codeudor.nombre AS codeudor_nombre,
                    codeudor.identificacion AS codeudor_identificacion,
                    COALESCE(p.tipo_cartera,'PREJURIDICO') AS tipo_cartera,
                    p.radicado_interno,
                    %s AS nombre_ph,
                    %s AS nit_ph
                FROM obligaciones o
                JOIN inmuebles_ph i ON i.id=o.inmueble_id
                JOIN LATERAL (
                    SELECT c.id AS contacto_id, c.nombre, c.identificacion
                    FROM obligacion_partes op
                    JOIN contactos c ON c.id=op.contacto_id
                    WHERE op.obligacion_id=o.id
                      AND op.rol='DEUDOR'
                    ORDER BY op.es_principal DESC, op.contacto_id ASC
                    LIMIT 1
                ) deudor ON TRUE
                LEFT JOIN LATERAL (
                    SELECT c.nombre, c.identificacion
                    FROM obligacion_partes op
                    JOIN contactos c ON c.id=op.contacto_id
                    WHERE op.obligacion_id=o.id
                      AND op.rol='CODEUDOR'
                    ORDER BY op.es_principal DESC, op.contacto_id ASC
                    LIMIT 1
                ) codeudor ON TRUE
                LEFT JOIN LATERAL (
                    SELECT p0.tipo_cartera, p0.radicado_interno
                    FROM procesos p0
                    JOIN proceso_obligaciones po0
                      ON po0.radicado_interno=p0.radicado_interno
                     AND po0.obligacion_id=o.id
                    ORDER BY po0.es_principal DESC, po0.id
                    LIMIT 1
                ) p ON TRUE
                WHERE {' AND '.join(where)}
                ORDER BY i.id, o.id DESC
                """,
                [
                    conjunto["nombre"],
                    conjunto.get("persona_juridica") or conjunto["nombre"],
                    conjunto.get("nit") or "",
                    *params,
                ],
            )
            filas = [dict(r) for r in cur.fetchall()]
    finally:
        if owns_conn and conn is not None:
            conn.release()

    corte = fecha_corte or date.today()
    minimo = float(saldo_minimo or 0)
    salida: list[dict] = []
    cache: dict[int, dict] = {}

    for fila in filas:
        oid = int(fila["obligacion_id"])
        if oid not in cache:
            cache[oid] = obligacion_saldo_service.calcular_saldo_obligacion(
                oid, fecha_corte=corte
            )
        saldo = cache[oid]
        if not saldo.get("saldo_verificado"):
            continue
        try:
            total = float(saldo.get("saldo_total") or 0)
        except (TypeError, ValueError):
            continue
        if total < minimo:
            continue

        item = dict(fila)
        item["saldo_total"] = round(total, 2)
        item["saldo_verificado"] = True
        item["saldo_fuente"] = saldo.get("saldo_fuente")
        item["fecha_corte"] = corte
        salida.append(item)

    return salida


def dedupe_por_inmueble(candidatos: list[dict]) -> list[dict]:
    """Garantiza una sola entrada por inmueble_id (primera gana)."""
    vistos: set[int] = set()
    salida: list[dict] = []
    for item in candidatos:
        iid = int(item["inmueble_id"])
        if iid in vistos:
            continue
        vistos.add(iid)
        salida.append(item)
    return salida


def resolver_seleccion(
    candidatos: list[dict],
    inmueble_ids: list[int],
) -> list[dict]:
    """Filtra por selección del usuario y vuelve a deduplicar."""
    ids = []
    vistos: set[int] = set()
    for raw in inmueble_ids:
        iid = int(raw)
        if iid in vistos:
            continue
        vistos.add(iid)
        ids.append(iid)

    if len(ids) > MAX_LOTE:
        raise ValueError(f"Máximo {MAX_LOTE} cartas por lote.")

    por_inmueble = {int(c["inmueble_id"]): c for c in dedupe_por_inmueble(candidatos)}
    seleccionados = []
    faltantes = []
    for iid in ids:
        item = por_inmueble.get(iid)
        if not item:
            faltantes.append(iid)
            continue
        seleccionados.append(item)
    if faltantes:
        raise ValueError(
            "Inmuebles no disponibles o sin saldo verificado: "
            + ", ".join(str(x) for x in faltantes[:10])
        )
    if not seleccionados:
        raise ValueError("Debe seleccionar al menos un inmueble.")
    return seleccionados


def construir_carta_docx(
    candidato: dict,
    *,
    fecha_limite: date,
    fecha_carta: Optional[date] = None,
) -> bytes:
    """Genera un .docx en memoria para un candidato."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(3)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    style.paragraph_format.space_after = Pt(0)

    fecha = fecha_carta or date.today()
    nombre = (candidato.get("deudor_nombre") or "SIN NOMBRE").strip()
    cedula = (candidato.get("deudor_identificacion") or "SIN CÉDULA").strip()
    torre = (candidato.get("torre_apto") or "SIN UNIDAD").strip()
    conjunto = (candidato.get("conjunto_residencial") or "SIN CONJUNTO").strip()
    nombre_ph = (candidato.get("nombre_ph") or conjunto).strip()
    monto = _fmt_money(candidato.get("saldo_total"))
    codeudor = (candidato.get("codeudor_nombre") or "").strip()

    def add_para(text: str = "", *, bold: bool = False, align=None, space_after: float = 6):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        return p

    add_para(f"{CIUDAD}, {_fmt_fecha_es(fecha)}.", space_after=18)
    add_para("Señor(a)", space_after=2)
    add_para(nombre, bold=True, space_after=2)
    add_para(f"C.C. {cedula}", space_after=2)
    add_para(
        f"Propietario(a) del ({torre}) ({conjunto}) {CIUDAD}.",
        space_after=6,
    )
    if codeudor:
        add_para(f"Atn. {codeudor}", space_after=12)
    else:
        add_para("", space_after=12)

    add_para(
        f"REF: REQUERIMIENTO DE PAGO PREJURÍDICO - {nombre_ph}",
        bold=True,
        space_after=14,
    )

    cuerpo = (
        f"Por medio de la presente, y obrando como apoderado de {nombre_ph}, "
        f"me permito requerirle el pago de la suma de {monto}, correspondiente "
        f"al total adeudado por concepto de cuotas de administración y demás "
        f"cargos asociados a la unidad ({torre}) del conjunto {conjunto}."
    )
    p_cuerpo = add_para(cuerpo, space_after=10)
    p_cuerpo.paragraph_format.first_line_indent = Cm(1.25)
    p_cuerpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    plazo = (
        f"Se le concede como fecha límite de pago el día {_fmt_fecha_es(fecha_limite)}. "
        "Vencido dicho plazo sin que se acredite el pago total, se adelantarán "
        "las gestiones prejurídicas y jurídicas a que haya lugar, con cargo de "
        "costas, intereses y honorarios."
    )
    p_plazo = add_para(plazo, space_after=10)
    p_plazo.paragraph_format.first_line_indent = Cm(1.25)
    p_plazo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    canales = (
        f"Para efectos de pago, información o radicación de acuerdos, puede "
        f"comunicarse al Tel/WhatsApp {TELEFONO_DESPACHO} o al correo "
        f"{CORREO_DESPACHO}."
    )
    p_canales = add_para(canales, space_after=18)
    p_canales.paragraph_format.first_line_indent = Cm(1.25)
    p_canales.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_para("Atentamente,", space_after=36)
    add_para(FIRMANTE_NOMBRE, bold=True, space_after=2)
    add_para(FIRMANTE_CARGO, space_after=2)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def nombre_archivo_carta(candidato: dict, fecha: Optional[date] = None) -> str:
    dia = (fecha or date.today()).strftime("%Y%m%d")
    conjunto = _safe_filename(candidato.get("conjunto_residencial") or "conjunto")
    unidad = _safe_filename(candidato.get("torre_apto") or "unidad")
    return f"Carta_Prejuridica_{conjunto}_{unidad}_{dia}.docx"


def generar_paquete_docx(
    seleccionados: list[dict],
    *,
    fecha_limite: date,
    fecha_carta: Optional[date] = None,
) -> tuple[bytes, str, str]:
    """Devuelve (contenido, filename, media_type). ZIP si hay más de una carta."""
    fecha = fecha_carta or date.today()
    if len(seleccionados) == 1:
        contenido = construir_carta_docx(
            seleccionados[0], fecha_limite=fecha_limite, fecha_carta=fecha
        )
        return contenido, nombre_archivo_carta(seleccionados[0], fecha), (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        usados: set[str] = set()
        for item in seleccionados:
            nombre = nombre_archivo_carta(item, fecha)
            if nombre in usados:
                stem = nombre.rsplit(".", 1)[0]
                nombre = f"{stem}_{item['inmueble_id']}.docx"
            usados.add(nombre)
            zf.writestr(
                nombre,
                construir_carta_docx(item, fecha_limite=fecha_limite, fecha_carta=fecha),
            )
    conjunto = _safe_filename(seleccionados[0].get("conjunto_residencial") or "conjunto")
    zip_name = f"Cartas_Prejuridicas_{conjunto}_{fecha.strftime('%Y%m%d')}.zip"
    return zip_buffer.getvalue(), zip_name, "application/zip"


def parse_fecha(value: Optional[str], *, default: date) -> date:
    texto = str(value or "").strip()
    if not texto:
        return default
    try:
        return datetime.strptime(texto, "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError("Fecha inválida. Use formato YYYY-MM-DD.") from exc
