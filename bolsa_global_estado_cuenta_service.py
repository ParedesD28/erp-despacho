"""
Integración Bolsa Global ↔ lote PDF estado de cuenta.

Separa la liquidación Art. 1653-1655 del Excel de solo-transcripción
(`generar_excel_lote`). Opera sobre movimientos ya parseados/sellados.
"""
from __future__ import annotations

import io
import json
from typing import Any, TypedDict

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bolsa_global_mora import determinar_mora_y_capital_limpio, resultado_vacio


class MetaCuentaBolsa(TypedDict):
    archivo: str
    titular: str
    bloque: str
    apartamento: str
    codigo_cuenta: str
    conjunto: str
    movimientos_extraidos: int
    error_parseo: str | None
    saldo_anterior_cuenta: float | None


def _meta_cuenta(cuenta: dict[str, Any]) -> MetaCuentaBolsa:
    saldo = cuenta.get("saldo_anterior")
    if saldo is None and isinstance(cuenta.get("cabecera"), dict):
        saldo = cuenta["cabecera"].get("saldo_anterior")
    return {
        "archivo": str(cuenta.get("archivo") or ""),
        "titular": str(cuenta.get("titular") or ""),
        "bloque": str(cuenta.get("bloque") or ""),
        "apartamento": str(cuenta.get("apartamento") or ""),
        "codigo_cuenta": str(cuenta.get("codigo_cuenta") or ""),
        "conjunto": str(cuenta.get("conjunto") or ""),
        "movimientos_extraidos": int(cuenta.get("movimientos_extraidos") or 0),
        "error_parseo": cuenta.get("error") or None,
        "saldo_anterior_cuenta": saldo if saldo is not None else None,
    }


def calcular_bolsa_por_cuenta(
    cuenta: dict[str, Any],
    *,
    extender_certificado_saldo_anterior: bool = False,
) -> dict[str, Any]:
    """Ejecuta Bolsa Global sobre una cuenta (un PDF / deudor)."""
    meta = _meta_cuenta(cuenta)
    if cuenta.get("error"):
        vacio = resultado_vacio(
            errores=[f"PDF no parseado: {cuenta.get('error')}"]
        )
        return {**meta, **vacio, "omitido": True}

    # Flag deprecado: se acepta por compatibilidad pero la Bolsa lo ignora.
    extender = bool(
        extender_certificado_saldo_anterior
        or cuenta.get("extender_certificado_saldo_anterior")
    )
    rows = cuenta.get("rows") or []
    resultado = determinar_mora_y_capital_limpio(
        rows,
        saldo_anterior=meta.get("saldo_anterior_cuenta"),
        extender_certificado_saldo_anterior=extender,
    )
    return {**meta, **resultado, "omitido": False}


def calcular_bolsa_lote(
    cuentas: list[dict[str, Any]] | None,
    *,
    extender_certificado_saldo_anterior: bool = False,
) -> dict[str, Any]:
    """
    Un bloque de resultado por deudor/archivo (sin mezclar rows entre cuentas).

    `cuentas` debe traer `rows` sellados (como en la caché del lote o el
    retorno interno de `procesar_lote_estados_cuenta` antes de strippear).
    `extender_certificado_saldo_anterior`: DEPRECADO, se ignora (el saldo
    anterior siempre se certifica en una fila «SALDO ANTERIOR»).
    """
    cuentas = list(cuentas or [])
    resultados = [
        calcular_bolsa_por_cuenta(
            c, extender_certificado_saldo_anterior=extender_certificado_saldo_anterior
        )
        for c in cuentas
    ]
    con_capital = 0
    total_lote = 0.0
    for r in resultados:
        total_lote += float(r.get("total_capital_demandado") or 0)
        if not r.get("omitido") and (r.get("capital_limpio_a_demandar") or []):
            con_capital += 1
    return {
        "cuentas_evaluadas": len(resultados),
        "cuentas_con_capital": con_capital,
        "total_capital_demandado_lote": round(total_lote, 2),
        "resultados": resultados,
    }


def resultado_a_json_bytes(payload: dict[str, Any]) -> io.BytesIO:
    raw = json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")
    return io.BytesIO(raw)


def generar_excel_bolsa_global(payload: dict[str, Any]) -> io.BytesIO:
    """
    Excel aparte del de transcripción: Resumen + Capital limpio + Errores.
    No modifica ni reutiliza `generar_excel_lote`.
    """
    resultados = list(payload.get("resultados") or [])
    resumen_rows = []
    capital_rows = []
    error_rows = []

    for r in resultados:
        resumen_rows.append(
            {
                "Archivo": r.get("archivo") or "",
                "Titular": r.get("titular") or "",
                "Bloque": r.get("bloque") or "",
                "Apartamento": r.get("apartamento") or "",
                "Codigo Cuenta": r.get("codigo_cuenta") or "",
                "Fecha inicio mora": r.get("fecha_inicio_mora") or "",
                "Total capital demandado": r.get("total_capital_demandado"),
                "Saldo anterior": r.get("saldo_anterior"),
                "Modo saldo anterior": r.get("saldo_anterior_modo") or "",
                "Alerta saldo anterior": bool(r.get("alerta_saldo_anterior")),
                "Saldo anterior pendiente": r.get("saldo_anterior_pendiente"),
                "Sanciones certificadas": r.get("sanciones_certificadas"),
                "Bolsa abonos (Σ)": r.get("bolsa_abonos"),
                "Bolsa inicial": r.get("bolsa_global_inicial"),
                "Bolsa remanente": r.get("bolsa_remanente_final"),
                "Ítems capital limpio": len(r.get("capital_limpio_a_demandar") or []),
                "Errores": len(r.get("errores_procesamiento") or []),
                "Omitido": bool(r.get("omitido")),
            }
        )
        for item in r.get("capital_limpio_a_demandar") or []:
            capital_rows.append(
                {
                    "Archivo": r.get("archivo") or "",
                    "Titular": r.get("titular") or "",
                    "Fecha": item.get("fecha") or "",
                    "Concepto": item.get("concepto") or "",
                    "Valor a demandar": item.get("valor_a_demandar"),
                    "Nota": item.get("nota") or "",
                    "Mes corte": item.get("mes_corte") or "",
                    "Clase": item.get("clase_concepto") or "",
                }
            )
        for err in r.get("errores_procesamiento") or []:
            error_rows.append(
                {
                    "Archivo": r.get("archivo") or "",
                    "Titular": r.get("titular") or "",
                    "Error": err,
                }
            )

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(resumen_rows).to_excel(writer, index=False, sheet_name="Resumen Bolsa")
        pd.DataFrame(capital_rows).to_excel(writer, index=False, sheet_name="Capital limpio")
        pd.DataFrame(error_rows).to_excel(writer, index=False, sheet_name="Errores")

        header_fill = PatternFill("solid", fgColor="0F172A")
        header_font = Font(color="FFFFFF", bold=True)
        wide_headers = frozenset({"Archivo", "Titular", "Concepto", "Nota", "Error"})
        for sheet_name in ("Resumen Bolsa", "Capital limpio", "Errores"):
            ws = writer.book[sheet_name]
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            for col in ws.columns:
                letter = get_column_letter(col[0].column)
                width = 36 if col[0].value in wide_headers else 18
                ws.column_dimensions[letter].width = width

    output.seek(0)
    return output
